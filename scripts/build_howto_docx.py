"""«Как пользоваться.docx» и «How to use.docx» из того же текста, что Справка в окне (static/i18n, ключи help.*) -
стандарт 5.1, описание 3.13.

    .venv312/bin/python scripts/build_howto_docx.py        # → docs/Как пользоваться.docx, docs/How to use.docx
"""
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent


class _Html(HTMLParser):
    """Подмножество разметки Справки: <p>, <ol>/<ul>/<li>, <b>, <span class="mono">."""

    def __init__(self, doc):
        super().__init__()
        self.doc, self.par, self.bold, self.mono, self.lists = doc, None, False, False, []

    def handle_starttag(self, tag, attrs):
        if tag == 'p':
            self.par = self.doc.add_paragraph()
        elif tag in ('ol', 'ul'):
            self.lists.append(tag)
        elif tag == 'li':
            self.par = self.doc.add_paragraph(style='List Number' if self.lists[-1] == 'ol' else 'List Bullet')
        elif tag == 'b':
            self.bold = True
        elif tag == 'span' and ('class', 'mono') in attrs:
            self.mono = True

    def handle_endtag(self, tag):
        if tag in ('p', 'li'):
            self.par = None
        elif tag in ('ol', 'ul'):
            self.lists.pop()
        elif tag == 'b':
            self.bold = False
        elif tag == 'span':
            self.mono = False

    def handle_data(self, data):
        if not data.strip() and self.par is None:
            return
        if self.par is None:
            self.par = self.doc.add_paragraph()
        run = self.par.add_run(data)
        run.bold = self.bold
        if self.mono:
            run.font.name = 'Menlo'


def build(lang: str, out: Path):
    t = json.loads((ROOT / 'static' / 'i18n' / f'{lang}.json').read_text(encoding='utf-8'))
    doc = Document()
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading(f"{t['app.name']} - {t['help.title']}", level=0)
    doc.add_paragraph(t['help.sub'])
    for i in range(1, 7):
        doc.add_heading(t[f'help.s{i}_t'], level=1)
        _Html(doc).feed(re.sub(r'\s+', ' ', t[f'help.s{i}']))
    doc.save(str(out))
    return out


if __name__ == '__main__':
    for lang, name in (('ru', 'Как пользоваться.docx'), ('en', 'How to use.docx')):
        print(build(lang, ROOT / 'docs' / name))
