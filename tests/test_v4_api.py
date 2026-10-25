"""Anonymizer 4.0 (задание 01): сервер по описанию 2.1 - сессии, Словарь, срок хранения, миграция, имена файлов.

Без окна и без моделей: Flask test client на своей временной папке данных (стандарт 6.10); документы вымышленные.
"""
import importlib
import io
import json
import re
import sqlite3
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tests' / 'bench'))
import ids  # noqa: E402

ids.R.seed(4001)
INN = ids.inn10()
TEXT = ('Договор поставки № 17. Поставщик - ООО «Северный ветер», ИНН {inn}, в лице генерального директора '
        'Петрова Ивана Сергеевича, действующего на основании Устава. Покупатель - Сидорова Анна Викторовна. '
        'Петров И. С. подписал договор.').format(inn=INN)


@pytest.fixture
def app_mod(tmp_path, monkeypatch):
    monkeypatch.setenv('ANONYMIZER_DATA_DIR', str(tmp_path / 'data'))
    monkeypatch.setenv('ANONYMIZER_NO_OPEN', '1')
    sys.modules.pop('app', None)
    mod = importlib.import_module('app')
    yield mod
    sys.modules.pop('app', None)


@pytest.fixture
def c(app_mod):
    return app_mod.app.test_client()


def _docx(text=TEXT) -> bytes:
    d = Document()
    for p in text.split('\n'):
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _session(c):
    r = c.post('/api/sessions', json={})
    assert r.status_code == 201
    return r.get_json()


def _upload(c, sid, name, data, section='anon'):
    return c.post(f'/api/sessions/{sid}/files', data={'file': (io.BytesIO(data), name), 'section': section},
                  content_type='multipart/form-data')


def _process(c, sid, name, section='anon', **kw):
    r = c.post(f'/api/sessions/{sid}/process', json={'input': name, 'section': section, 'use_spacy': False, **kw})
    assert r.status_code == 200, r.get_json()
    return r.get_json()


def _anon_one(c, name='Договор поставки - Петров.docx', text=TEXT):
    s = _session(c)
    assert _upload(c, s['id'], name, _docx(text)).status_code == 201
    return s, _process(c, s['id'], name)


def _values(c, sid):
    return {m['original_form'] for m in c.get(f'/api/sessions/{sid}/mappings?all=1').get_json()}


# ── каркас ───────────────────────────────────────────────────────────────────
def test_info_version_matches_window(c):
    """Стандарт 6.5: /api/info = core/version = UI_VERSION окна; порт 8070 (6.3)."""
    from core.version import __version__
    info = c.get('/api/info').get_json()
    assert info['version'] == __version__ == '4.0.1'
    js = (ROOT / 'static' / 'js' / 'core.js').read_text(encoding='utf-8')
    assert f"const UI_VERSION = '{__version__}'" in js


def test_default_port_and_data_dir_outside_project(app_mod, tmp_path):
    assert app_mod.PORT in (8070, 8074) or 'ANONYMIZER_PORT' in __import__('os').environ
    assert app_mod.DATA_DIR == tmp_path / 'data'
    src = (ROOT / 'app.py').read_text(encoding='utf-8')
    assert "or 8070" in src and 'Path(__file__).parent\n' not in src


def test_dead_api_removed(c):
    """5.5: /api/patterns без кнопки - удалён."""
    assert c.get('/api/patterns').status_code == 404


# ── сессии ───────────────────────────────────────────────────────────────────
def test_session_name_never_from_file_name(c):
    """3.9, было Р-39: «Сессия 24.10», вторая за день - «(2)»; имя файла в имя сессии не идёт."""
    s, r = _anon_one(c)
    assert re.fullmatch(r'Сессия \d\d\.\d\d', s['name'])
    s2 = _session(c)
    assert s2['name'] == s['name'] + ' (2)'
    names = [x['name'] for x in c.get('/api/sessions').get_json()]
    assert not any('Петров' in n or 'Договор' in n for n in names)
    row = next(x for x in c.get('/api/sessions').get_json() if x['id'] == s['id'])
    assert row['files'] == 1 and row['records'] > 0 and row['expires_in_days'] == 30


def test_records_count_is_one_number(c):
    """Было Р-16: число записей в «Сессиях» = активные метки таблицы замен."""
    s, r = _anon_one(c)
    tokens = {m['token'] for m in c.get(f"/api/sessions/{s['id']}/mappings").get_json()}
    row = next(x for x in c.get('/api/sessions').get_json() if x['id'] == s['id'])
    assert row['records'] == len(tokens) == r['records']


def test_rename_session(c):
    s = _session(c)
    assert c.patch(f"/api/sessions/{s['id']}", json={'name': 'Аренда офиса'}).status_code == 200
    assert c.get(f"/api/sessions/{s['id']}").get_json()['name'] == 'Аренда офиса'


def test_delete_session_whole_and_dictionary_stays(c, app_mod):
    """3.10, UC-S5, UC-W3: папки нет, записей в базе нет; Словарь не тронут."""
    s, _ = _anon_one(c)
    c.post('/api/dictionary', json={'value': 'Ковалёв Дмитрий', 'entity_type': 'ФИО', 'list': 'always'})
    assert (app_mod.UPLOADS_DIR / s['id']).exists()
    assert c.delete(f"/api/sessions/{s['id']}").status_code == 200
    assert not (app_mod.UPLOADS_DIR / s['id']).exists()
    con = sqlite3.connect(str(app_mod.DB_PATH))
    for table in ('mappings', 'places', 'place_forms', 'exclusions', 'occurrences'):
        assert con.execute(f'SELECT COUNT(*) FROM {table} WHERE session_id=?', (s['id'],)).fetchone()[0] == 0, table
    con.close()
    assert [x['value'] for x in c.get('/api/dictionary').get_json()['always']] == ['Ковалёв Дмитрий']


def test_retention_deletes_whole_session(c, app_mod):
    """3.9, UC-S3: старше срока - сессия целиком (записи, места, папка); за 3 дня - «удалится через 3 дня»."""
    old, _ = _anon_one(c)
    near = _session(c)
    con = sqlite3.connect(str(app_mod.DB_PATH))
    con.execute("UPDATE sessions SET created_at = datetime('now','localtime','-31 days') WHERE id=?", (old['id'],))
    con.execute("UPDATE sessions SET created_at = datetime('now','localtime','-27 days','-1 hours') WHERE id=?", (near['id'],))
    con.commit()
    con.close()
    rows = {x['id']: x for x in c.get('/api/sessions').get_json()}
    assert rows[near['id']]['expires_in_days'] == 3
    assert app_mod.purge_expired() == 1
    assert old['id'] not in {x['id'] for x in c.get('/api/sessions').get_json()}
    assert not (app_mod.UPLOADS_DIR / old['id']).exists()
    con = sqlite3.connect(str(app_mod.DB_PATH))
    assert con.execute('SELECT COUNT(*) FROM mappings WHERE session_id=?', (old['id'],)).fetchone()[0] == 0
    con.close()


def test_retention_setting_only_7_30_90(c):
    """UC-N1, вопрос 3: «не удалять» нет."""
    assert c.post('/api/settings', json={'retention_days': 7}).get_json()['retention_days'] == 7
    assert c.post('/api/settings', json={'retention_days': 0}).get_json()['retention_days'] == 7
    assert c.post('/api/settings', json={'retention_days': 365}).get_json()['retention_days'] == 7
    assert c.get('/api/settings').get_json()['retention_days'] == 7


# ── файлы ────────────────────────────────────────────────────────────────────
def test_download_names_only_with_tokens(c):
    """Правка 6: в окне - исходное имя, в скачиваемом файле и ZIP - только метки, имя ZIP без текста пользователя."""
    s, r = _anon_one(c)
    out = r['file']['output']
    assert r['file']['input'] == 'Договор поставки - Петров.docx'
    assert 'Петров' not in out and re.search(r'\[FIO_\d+\]', out), out
    _upload(c, s['id'], 'Акт - Сидорова.docx', _docx(TEXT + '\nАкт подписала Сидорова Анна Викторовна.'))
    _process(c, s['id'], 'Акт - Сидорова.docx')
    z = c.get(f"/api/download-all/{s['id']}?section=anon")
    assert z.status_code == 200
    cd = z.headers['Content-Disposition']
    assert 'Петров' not in cd and 'Сидоров' not in cd
    names = zipfile.ZipFile(io.BytesIO(z.data)).namelist()
    values = _values(c, s['id'])
    assert len(names) == 2
    for n in names:
        assert not any(v in n for v in values if len(v) > 3), n


def test_zip_only_its_section(c):
    """Было Р-40: ZIP «Обезличить» без файлов «Вернуть данные»."""
    s, r = _anon_one(c)
    anon_bytes = c.get(f"/api/download/{s['id']}/{r['file']['output']}").data
    assert _upload(c, s['id'], r['file']['output'], anon_bytes, 'restore').status_code == 201
    rr = _process(c, s['id'], r['file']['output'], 'restore')
    assert rr['file']['output'].endswith(' (данные возвращены).docx')
    anon = zipfile.ZipFile(io.BytesIO(c.get(f"/api/download-all/{s['id']}?section=anon").data)).namelist()
    back = c.get(f"/api/download-all/{s['id']}?section=restore")
    assert anon == [r['file']['output']]
    assert zipfile.ZipFile(io.BytesIO(back.data)).namelist() == [rr['file']['output']]


def test_restored_file_name(c):
    """Правка 1, UC-D5: «Претензия покупателю (данные возвращены).docx», не `_restored`."""
    s = _session(c)
    _upload(c, s['id'], 'Претензия покупателю.docx', _docx(TEXT))
    r = _process(c, s['id'], 'Претензия покупателю.docx')
    data = c.get(f"/api/download/{s['id']}/{r['file']['output']}").data
    _upload(c, s['id'], 'Претензия покупателю.docx', data, 'restore')
    rr = _process(c, s['id'], 'Претензия покупателю.docx', 'restore')
    assert rr['file']['output'] == 'Претензия покупателю (данные возвращены).docx'


def test_sections_survive_reopen(c):
    """UC-A5: раздел и итоги - из manifest.json, не по имени файла."""
    s, r = _anon_one(c)
    data = c.get(f"/api/download/{s['id']}/{r['file']['output']}").data
    _upload(c, s['id'], 'Ответ нейросети.docx', data, 'restore')
    _process(c, s['id'], 'Ответ нейросети.docx', 'restore')
    files = c.get(f"/api/sessions/{s['id']}").get_json()['file_list']
    assert [(f['input'], f['section'], f['status']) for f in files] == [
        ('Договор поставки - Петров.docx', 'anon', 'done'), ('Ответ нейросети.docx', 'restore', 'done')]
    assert files[0]['result']['found'] > 0 and files[0]['result']['leaks'] == []
    assert files[1]['result']['restore']['restored'] > 0


def test_duplicate_and_format_and_size(c):
    """4.6: тот же файл второй раз - «Этот файл уже в сессии», чужой формат - код format."""
    s = _session(c)
    data = _docx()
    assert _upload(c, s['id'], 'a.docx', data).status_code == 201
    r = _upload(c, s['id'], 'a.docx', data)
    assert r.status_code == 409 and r.get_json()['code'] == 'duplicate'
    r = _upload(c, s['id'], 'a.exe', b'MZ')
    assert r.status_code == 400 and r.get_json()['code'] == 'format'


def test_broken_file_error_code(c):
    s = _session(c)
    _upload(c, s['id'], 'broken.docx', b'PK\x03\x04 not really a zip')
    r = _process(c, s['id'], 'broken.docx')
    assert r['file']['status'] == 'error' and r['file']['error'] == 'broken'


def test_remove_file_keeps_records(c, app_mod):
    """UC-A7: файла нет в списке и на диске, записи таблицы замен на месте."""
    s, r = _anon_one(c)
    before = _values(c, s['id'])
    assert c.delete(f"/api/sessions/{s['id']}/files", json={'input': 'Договор поставки - Петров.docx', 'section': 'anon'}).status_code == 200
    assert c.get(f"/api/sessions/{s['id']}/files").get_json()['files'] == []
    assert not list((app_mod.UPLOADS_DIR / s['id'] / 'input').iterdir())
    assert not list((app_mod.UPLOADS_DIR / s['id'] / 'output').iterdir())
    assert _values(c, s['id']) == before


def test_legacy_manifest_read():
    from core import session_files as sf
    import tempfile
    d = Path(tempfile.mkdtemp())
    (d / 'manifest.json').write_text(json.dumps({'a.docx': {'anonymize': 'a [FIO_1].docx', 'anonymize_pdf': 'a.pdf'},
                                                 'b.docx': {'deanonymize': 'b_restored.docx'}}), encoding='utf-8')
    files = sf.load(d)
    assert [(f['input'], f['section'], f['output'], f['extra']) for f in files] == [
        ('a.docx', 'anon', 'a [FIO_1].docx', 'a.pdf'), ('b.docx', 'restore', 'b_restored.docx', None)]


# ── таблица замен и Словарь ──────────────────────────────────────────────────
def test_unmask_and_mask_again_same_token(c):
    """UC-M3: «Не маскировать» → «Не маскируется» → «Вернуть маску» → та же метка в файле."""
    s, r = _anon_one(c)
    sid = s['id']
    maps = c.get(f'/api/sessions/{sid}/mappings').get_json()
    tok = next(m['token'] for m in maps if m['entity_type'] == 'ФИО' and 'Петров' in m['original_form'])
    assert c.delete(f'/api/sessions/{sid}/mappings/{tok}', json={'scope': 'session'}).status_code == 200
    assert all(m['token'] != tok for m in c.get(f'/api/sessions/{sid}/mappings').get_json())
    excluded = [m for m in c.get(f'/api/sessions/{sid}/mappings?all=1').get_json() if m['status'] == 'excluded']
    assert {m['token'] for m in excluded} == {tok}
    files = c.post(f'/api/sessions/{sid}/reprocess', json={}).get_json()['results']
    text = '\n'.join(p.text for p in Document(io.BytesIO(c.get(f"/api/download/{sid}/{files[0]['output']}").data)).paragraphs)
    assert 'Петров' in text and f'[{tok}]' not in text
    assert c.post(f'/api/sessions/{sid}/mappings/{tok}/mask', json={}).status_code == 200
    files = c.post(f'/api/sessions/{sid}/reprocess', json={}).get_json()['results']
    text = '\n'.join(p.text for p in Document(io.BytesIO(c.get(f"/api/download/{sid}/{files[0]['output']}").data)).paragraphs)
    assert 'Петров' not in text and f'[{tok}]' in text
    assert files[0]['leaks'] == []


def test_unmask_always_goes_to_never_list_and_works_in_new_session(c):
    """UC-U1, 3.11: «всегда» - в «Никогда не маскировать», в новой сессии значение не маскируется."""
    s, _ = _anon_one(c)
    maps = c.get(f"/api/sessions/{s['id']}/mappings").get_json()
    m = next(x for x in maps if x['entity_type'] == 'ФИО' and 'Сидорова' in x['original_form'])
    c.delete(f"/api/sessions/{s['id']}/mappings/{m['token']}", json={'scope': 'always'})
    d = c.get('/api/dictionary').get_json()
    assert [(x['value'], x['source']) for x in d['never']] == [(m['original_form'], 'unmask')]
    s2, _ = _anon_one(c, 'Второй.docx')
    assert not any('Сидорова' in v for v in _values(c, s2['id']))


def test_manual_add_only_always_goes_to_dictionary(c):
    """Было Р-24: в Словарь - только при «всегда»; UC-W1: «всегда» маскирует в новой сессии."""
    s = _session(c)
    sid = s['id']
    r = c.post(f'/api/sessions/{sid}/mappings', json={'original_form': 'Голубев', 'entity_type': 'ФИО'})
    assert r.status_code == 201
    assert c.get('/api/dictionary').get_json()['always'] == []
    r = c.post(f'/api/sessions/{sid}/mappings', json={'original_form': 'Альфа-Трейд', 'entity_type': 'ЮЛ', 'scope': 'always'})
    assert r.status_code == 201
    assert [x['value'] for x in c.get('/api/dictionary').get_json()['always']] == ['Альфа-Трейд']
    s2, _ = _anon_one(c, 'Новый.docx', TEXT + '\nСоисполнитель - Альфа-Трейд.')
    out = c.get(f"/api/sessions/{s2['id']}/files").get_json()['files'][0]['output']
    text = '\n'.join(p.text for p in Document(io.BytesIO(c.get(f"/api/download/{s2['id']}/{out}").data)).paragraphs)
    assert 'Альфа-Трейд' not in text


def test_manual_add_before_first_file(c):
    """UC-M2: запись до первого файла применяется при обработке."""
    s = _session(c)
    c.post(f"/api/sessions/{s['id']}/mappings", json={'original_form': 'Северный причал', 'entity_type': 'ЮЛ'})
    _upload(c, s['id'], 'x.docx', _docx(TEXT + '\nСклад - Северный причал.'))
    r = _process(c, s['id'], 'x.docx')
    text = '\n'.join(p.text for p in Document(io.BytesIO(c.get(f"/api/download/{s['id']}/{r['file']['output']}").data)).paragraphs)
    assert 'Северный причал' not in text


def test_dictionary_lists_exclusive_and_undo(c):
    """3.11: значение не бывает в двух списках; ✕ и «Вернуть» с тем же источником и датой."""
    assert c.post('/api/dictionary', json={'value': 'Иванов', 'entity_type': 'ФИО', 'list': 'always'}).get_json()['moved'] is False
    assert c.post('/api/dictionary', json={'value': 'Иванов', 'entity_type': 'ФИО', 'list': 'never'}).get_json()['moved'] is True
    d = c.get('/api/dictionary').get_json()
    assert d['always'] == [] and [x['value'] for x in d['never']] == ['Иванов']
    row = c.post('/api/dictionary/remove', json={'value': 'Иванов', 'entity_type': 'ФИО', 'list': 'never'}).get_json()['row']
    assert c.get('/api/dictionary').get_json()['never'] == []
    c.post('/api/dictionary', json={**row, 'list': 'never'})
    back = c.get('/api/dictionary').get_json()['never'][0]
    assert back['added_at'] == row['added_at'] and back['source'] == row['source']
    assert c.post('/api/dictionary/clear', json={'list': 'never'}).get_json()['removed'] == 1


def test_base_form_nominative(c):
    """Правка 2: «Начальная форма» ФИО - в именительном падеже."""
    s, _ = _anon_one(c)
    maps = c.get(f"/api/sessions/{s['id']}/mappings").get_json()
    p = next(m for m in maps if m['original_form'].startswith('Петрова Ивана'))
    assert p['base_form'] == 'Петров Иван Сергеевич'


# ── возврат данных ───────────────────────────────────────────────────────────
def test_restore_pasted_text(c):
    """Вопрос 1, UC-D4: вставленный ответ - падеж подобран, чужая метка названа, на диск ничего не пишется."""
    s, _ = _anon_one(c)
    sid = s['id']
    maps = c.get(f'/api/sessions/{sid}/mappings').get_json()
    fio = next(m['token'] for m in maps if m['original_form'].startswith('Петрова Ивана'))
    r = c.post(f'/api/sessions/{sid}/restore-text', json={'text': f'Направляем [{fio}] ответ.\nКопия - [FIO_99].'})
    d = r.get_json()
    assert 'Петрову Ивану Сергеевичу' in d['text'] and '[FIO_99]' in d['text']
    marks = {d['text'][a:b]: m for a, b, m in d['spans']}
    assert marks['Петрову Ивану Сергеевичу'] == 'case' and marks['[FIO_99]'] == 'unknown'
    assert d['stats']['unknown'] == ['FIO_99']
    files = c.get(f'/api/sessions/{sid}/files').get_json()['files']
    assert [f['section'] for f in files] == ['anon']


def test_changed_initial_form_used_in_new_text(c):
    """UC-D1: изменённая «Начальная форма» ФИО склоняется в новом тексте."""
    s, _ = _anon_one(c)
    sid = s['id']
    maps = c.get(f'/api/sessions/{sid}/mappings').get_json()
    m = next(x for x in maps if x['original_form'].startswith('Петрова Ивана'))
    c.patch(f"/api/sessions/{sid}/mappings/{m['token']}", json={'canonical_form': 'Петров Илья Сергеевич'})
    d = c.post(f'/api/sessions/{sid}/restore-text', json={'text': f"Направляем [{m['token']}] ответ."}).get_json()
    assert 'Петрову Илье Сергеевичу' in d['text'], d['text']


def test_restore_without_llm_setting_does_not_call_llm(c, monkeypatch):
    """Вопрос 2: нейросеть выключена - падеж только правилами (спорное место подсвечено), запрос к ней не уходит."""
    from core import llm
    called = []
    monkeypatch.setattr(llm, 'choose_cases', lambda *a, **k: called.append(1) or {})
    s, _ = _anon_one(c)
    maps = c.get(f"/api/sessions/{s['id']}/mappings").get_json()
    fio = next(m['token'] for m in maps if m['entity_type'] == 'ФИО')
    c.post(f"/api/sessions/{s['id']}/restore-text", json={'text': f'[{fio}]'})
    assert not called


# ── миграция 4.0 (5.4) ───────────────────────────────────────────────────────
def _legacy_db(path: Path):
    con = sqlite3.connect(str(path))
    con.executescript('''
        CREATE TABLE sessions (id TEXT PRIMARY KEY, name TEXT NOT NULL, created_at TEXT DEFAULT (datetime('now','localtime')));
        CREATE TABLE mappings (id INTEGER PRIMARY KEY AUTOINCREMENT, session_id TEXT NOT NULL, token TEXT NOT NULL,
            original_form TEXT NOT NULL DEFAULT '', canonical_form TEXT NOT NULL, entity_type TEXT NOT NULL,
            created_at TEXT DEFAULT (datetime('now','localtime')), canonical_edited INTEGER NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'active', opf TEXT, UNIQUE(session_id, original_form, entity_type));
        CREATE TABLE user_patterns (id INTEGER PRIMARY KEY AUTOINCREMENT, pattern TEXT NOT NULL, entity_type TEXT NOT NULL,
            created_at TEXT);
        CREATE TABLE exclusions (session_id TEXT NOT NULL, original_form TEXT NOT NULL, entity_type TEXT NOT NULL,
            PRIMARY KEY (session_id, original_form, entity_type));
        CREATE TABLE known_entities (value TEXT NOT NULL, entity_type TEXT NOT NULL, seen_count INTEGER NOT NULL DEFAULT 1,
            last_seen_at TEXT DEFAULT (datetime('now','localtime')), PRIMARY KEY (value, entity_type));
        INSERT INTO sessions (id, name) VALUES ('s1', 'Договор · 01.10.26 10:00');
        INSERT INTO mappings (session_id, token, original_form, canonical_form, entity_type) VALUES
            ('s1', 'FIO_1', 'Котов Олег', 'Котов Олег', 'ФИО'), ('s1', 'INN_1', '7707083893', '7707083893', 'ИНН');
        INSERT INTO known_entities (value, entity_type) VALUES ('Котов Олег', 'ФИО'), ('Ромашка', 'ЮЛ');
        INSERT INTO user_patterns (pattern, entity_type) VALUES ('в лице [ФИО]', 'ФИО');
        PRAGMA user_version = 4;
    ''')
    con.commit()
    con.close()


def test_migration_backup_counts_and_dictionary(tmp_path):
    from core import migrate
    from core.db import init_db, dict_list
    _legacy_db(tmp_path / 'anon.db')
    (tmp_path / 'settings.json').write_text('{"org_opf": "never"}', encoding='utf-8')
    r = migrate.migrate(tmp_path, init_db)
    assert r['done'] and r['moved'] == 2 and all(r['after'][k] == v for k, v in r['before'].items())
    assert (tmp_path / 'backup-before-4.0' / 'anon.db').exists()
    assert (tmp_path / 'backup-before-4.0' / 'settings.json').read_text(encoding='utf-8') == '{"org_opf": "never"}'
    assert migrate.counts(tmp_path / 'backup-before-4.0' / 'anon.db') == r['before']
    assert {x['source'] for x in dict_list(tmp_path / 'anon.db', 'always')} == {'legacy'}
    assert migrate.migrate(tmp_path, init_db) == {'done': False}


def test_migration_rolls_back_on_mismatch(tmp_path):
    from core import migrate
    from core.db import init_db
    _legacy_db(tmp_path / 'anon.db')

    def broken(db):
        init_db(db)
        con = sqlite3.connect(str(db))
        con.execute('DELETE FROM mappings')
        con.commit()
        con.close()
    r = migrate.migrate(tmp_path, broken)
    assert 'error' in r
    assert migrate.counts(tmp_path / 'anon.db')['mappings'] == 2


def test_migrated_count_shown_once(tmp_path, monkeypatch):
    """Правка 5: «Перенесено N значений из прежней памяти» - до ✕."""
    data = tmp_path / 'data'
    data.mkdir()
    _legacy_db(data / 'anon.db')
    monkeypatch.setenv('ANONYMIZER_DATA_DIR', str(data))
    sys.modules.pop('app', None)
    mod = importlib.import_module('app')
    c = mod.app.test_client()
    assert c.get('/api/dictionary').get_json()['migrated'] == 2
    assert c.get('/api/info').get_json()['migration'] == {'done': True, 'moved': 2}
    c.post('/api/dictionary/note-hide', json={})
    assert c.get('/api/dictionary').get_json()['migrated'] == 0
    sys.modules.pop('app', None)
