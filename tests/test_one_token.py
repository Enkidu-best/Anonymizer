"""One entity = one token in every form; exact forms come back (CLAUDE_CODE_TASK.md stage 5)."""
from docx import Document

from core.anonymizer import anonymize_text, restore_text
from core.handlers import process_uploaded_file

TEXT = ('Гражданин Белозёров Аркадий Львович (далее — Продавец). '
        'Белозёрова Аркадия Львовича уведомили. Белозёрову Аркадию Львовичу передано. '
        'С Белозёровым Аркадием Львовичем согласовано. О Белозёрове Аркадии Львовиче известно. '
        'Белозёров обязуется передать долю. Также ООО «Вектор-Плюс» и Вектор-Плюс подтверждают. '
        'Подпись: А.Л. Белозёров')
FORMS = ['Белозёров Аркадий Львович', 'Белозёрова Аркадия Львовича', 'Белозёрову Аркадию Львовичу',
         'Белозёровым Аркадием Львовичем', 'Белозёрове Аркадии Львовиче', 'Белозёров обязуется',
         'А.Л. Белозёров']


def test_all_forms_one_token_and_exact_restore(tmp_db, session_id):
    out, occ = anonymize_text(TEXT, tmp_db, session_id, use_spacy=False)
    assert 'Белозёр' not in out, out
    assert 'Вектор-Плюс' not in out, out
    fio = [t for t in occ if t.startswith('FIO')]
    assert len(fio) == 1, occ
    assert len(occ[fio[0]]) == 7
    yul = [t for t in occ if t.startswith('YUL')]
    assert len(yul) == 1, occ
    assert restore_text(out, tmp_db, session_id, occ) == TEXT


def test_llm_reply_gets_main_form(tmp_db, session_id):
    out, occ = anonymize_text(TEXT, tmp_db, session_id, use_spacy=False)
    tok = next(t for t in occ if t.startswith('FIO'))
    reply = f'Риск: [{tok}] может не передать долю.'
    assert restore_text(reply, tmp_db, session_id) == 'Риск: Белозёров Аркадий Львович может не передать долю.'


def test_docx_file_forms_restored(tmp_path, tmp_db, session_id):
    src = tmp_path / 'Договор_Белозёров.docx'
    doc = Document()
    for sentence in TEXT.split('. '):
        doc.add_paragraph(sentence)
    doc.save(src)
    (tmp_path / 'a').mkdir()
    (tmp_path / 'b').mkdir()
    r = process_uploaded_file(src, tmp_path / 'a', session_id, tmp_db, 'anonymize', use_spacy=False)
    assert 'Белозёр' not in r['output_filename']
    anon = tmp_path / 'a' / r['output_filename']
    assert 'Белозёр' not in '\n'.join(p.text for p in Document(anon).paragraphs)
    r2 = process_uploaded_file(anon, tmp_path / 'b', session_id, tmp_db, 'deanonymize')
    assert r2['output_filename'].startswith('Договор Белозёров')
    back = '\n'.join(p.text for p in Document(tmp_path / 'b' / r2['output_filename']).paragraphs)
    assert back == '\n'.join(TEXT.split('. '))
