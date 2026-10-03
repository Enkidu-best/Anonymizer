"""Restoring a completely different text with the same placeholders (task 3, §2.5).

The generator document is anonymized; an «external LLM» reply (tests/fixtures/ai_rewrite.py)
is restored: no tokens left, names in the right case, companies in quotes with the legal
form where needed, inserted value keeps the font size of the token's run, places with
the case chosen are highlighted, the journal holds no values.
"""
import os
from pathlib import Path

from docx import Document

from core.db import get_session_mappings
from core.handlers import process_uploaded_file
from tests.fixtures import ai_rewrite, generator


def _prepare(tmp_path, db, sid):
    src = tmp_path / 'contract.docx'
    generator.make_docx(src)
    (tmp_path / 'a').mkdir()
    process_uploaded_file(src, tmp_path / 'a', sid, db, 'anonymize', use_spacy=False)
    return ai_rewrite.make(tmp_path, get_session_mappings(db, sid))


def _paragraphs(doc):
    paras = [p for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            paras.append(row.cells[1].paragraphs[0])
    return paras


def test_foreign_text_docx(tmp_path, tmp_db, session_id):
    reply = _prepare(tmp_path, tmp_db, session_id)
    assert set(reply['tokens']) >= {'T', 'R', 'ORG', 'INN', 'ADR'}, reply['tokens']
    (tmp_path / 'b').mkdir()
    r = process_uploaded_file(reply['docx'], tmp_path / 'b', session_id, tmp_db, 'deanonymize')
    back = Document(tmp_path / 'b' / r['output_filename'])
    paras = _paragraphs(back)
    expected = [b for b in reply['blocks']]
    assert [p.text for p in paras] == [b[3] for b in expected]
    for p, (kind, size, _, want, names) in zip(paras, expected):
        assert '[' not in p.text and '_' not in p.text, p.text
        for run in p.runs:                 # every inserted piece keeps the token run's size
            if run.text.strip():
                assert run.font.size.pt == size, (kind, run.text, run.font.size.pt)
        hl = [x.text for x in p.runs if x.font.highlight_color is not None]
        assert hl == names, (kind, hl)
    assert r['restore']['unknown'] == []
    assert r['restore']['check_case'] == sum(len(b[4]) for b in expected)


def test_foreign_text_txt(tmp_path, tmp_db, session_id):
    reply = _prepare(tmp_path, tmp_db, session_id)
    (tmp_path / 'b').mkdir()
    r = process_uploaded_file(reply['txt'], tmp_path / 'b', session_id, tmp_db, 'deanonymize')
    back = (tmp_path / 'b' / r['output_filename']).read_text(encoding='utf-8').split('\n')
    assert back == [b[3] for b in reply['blocks']]


def test_journal_has_no_values(tmp_path, tmp_db, session_id):
    reply = _prepare(tmp_path, tmp_db, session_id)
    (tmp_path / 'b').mkdir()
    process_uploaded_file(reply['docx'], tmp_path / 'b', session_id, tmp_db, 'deanonymize')
    logs = Path(os.environ['ANONYMIZER_LOG_DIR'])
    text = ''.join(p.read_text(encoding='utf-8', errors='replace') for p in logs.rglob('*') if p.is_file())
    for v in ('Тарханов', 'Ракитин', 'Северная', reply['tokens']['INN_VALUE']):
        assert v not in text
