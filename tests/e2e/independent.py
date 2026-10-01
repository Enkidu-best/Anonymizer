"""Independent text extraction for the end-to-end test.

Deliberately does NOT use core.ooxml / core.rtf / core.extract (the code that writes
the files): a bug in the writer must not hide itself in the check. Returns «places»
— paragraphs, cells, pages — so differences can be reported by place number.
"""
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from urllib.parse import unquote

from lxml import etree

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
_P = {f'{{{W}}}p', f'{{{A}}}p'}
_T = {f'{{{W}}}t', f'{{{A}}}t', f'{{{W}}}instrText'}
_SKIP_ANC = {f'{{{W}}}del', f'{{{W}}}moveFrom'}   # deleted revision text is not in the document


def places(path) -> dict:
    """{'body': [...], 'meta': [...]} — body = text places, meta = properties, link targets."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in ('.docx', '.docm', '.pptx'):
        return _ooxml(path)
    if ext in ('.xlsx', '.xlsm'):
        return _xlsx(path)
    if ext == '.pdf':
        return _pdf(path)
    if ext == '.rtf':
        return {'body': _rtf(path), 'meta': []}
    if ext in ('.txt', '.md', '.csv'):
        return {'body': path.read_text(encoding='utf-8', errors='replace').split('\n'), 'meta': []}
    if ext in ('.jpg', '.jpeg', '.png', '.tif', '.tiff', '.heic', '.bmp', '.webp'):
        return {'body': [], 'meta': [], 'ocr': _ocr_image(path)}
    raise ValueError(ext)


def _ooxml(path):
    body, meta = [], []
    with zipfile.ZipFile(path) as z:
        for name in sorted(z.namelist()):
            data = z.read(name)
            if name.endswith('.rels'):
                for m in re.finditer(rb'Target="([^"]+)"[^>]*TargetMode="External"|TargetMode="External"[^>]*Target="([^"]+)"', data):
                    meta.append(unquote((m.group(1) or m.group(2)).decode('utf-8', 'replace')))
                continue
            if not name.endswith('.xml'):
                continue
            try:
                root = etree.fromstring(data)
            except etree.XMLSyntaxError:
                continue
            if name.startswith('docProps/'):
                meta += [e.text for e in root.iter() if isinstance(e.tag, str) and len(e) == 0 and e.text]
                continue
            groups = {}
            for el in root.iter():
                if el.tag not in _T and el.tag not in (f'{{{W}}}tab', f'{{{W}}}br', f'{{{A}}}br'):
                    continue
                p, skip = None, False
                for anc in el.iterancestors():
                    if anc.tag in _SKIP_ANC:
                        skip = True
                    if anc.tag in _P:
                        p = anc
                        break
                if p is None or skip:
                    continue
                piece = (el.text or '') if el.tag in _T else ('\t' if el.tag.endswith('}tab') else '\n')
                groups.setdefault(id(p), [p, []])[1].append(piece)
            for p in root.iter(*_P):
                if id(p) in groups:
                    body.append(''.join(groups[id(p)][1]))
    return {'body': body, 'meta': meta}


def _xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    body, meta = [], []
    for ws in wb.worksheets:
        body.append(ws.title)
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip() != '':
                    body.append(str(c.value))
                if c.comment:
                    body.append(c.comment.text)
    p = wb.properties
    meta += [x for x in (p.creator, p.lastModifiedBy, p.title, p.subject, p.description) if x]
    return {'body': body, 'meta': meta}


def _pdf(path):
    import pymupdf
    doc = pymupdf.open(str(path))
    body, ocr = [], []
    for page in doc:
        body.append(page.get_text())
        if page.get_images():   # what a reader SEES on a scan — checked for leaked values
            ocr.append('\n'.join(_ocr_pixmap(page.get_pixmap(dpi=150))))
    meta = [v for v in (doc.metadata or {}).values() if v] + [t[1] for t in doc.get_toc()]
    return {'body': body, 'meta': meta, 'ocr': ocr}


def _rtf(path):
    if sys.platform == 'darwin' and shutil.which('textutil'):
        out = Path(tempfile.mkdtemp()) / 'x.txt'
        subprocess.run(['textutil', '-convert', 'txt', '-output', str(out), str(path)],
                       check=True, capture_output=True)
        return out.read_text(encoding='utf-8', errors='replace').split('\n')
    from striprtf.striprtf import rtf_to_text
    return rtf_to_text(path.read_text(encoding='latin-1'), encoding='cp1251', errors='replace').split('\n')


def _ocr_pixmap(pix):
    try:
        from core import ocr   # OCR engine is shared — there is only one on the machine
        if not ocr.available():
            return []
        return [l.text for l in ocr.recognize(pix.tobytes('png'), (pix.width, pix.height))]
    except Exception:
        return []


def _ocr_image(path):
    import pymupdf
    return _ocr_pixmap(pymupdf.Pixmap(str(path)))


def norm(s: str) -> str:
    return re.sub(r'\s+', ' ', (s or '').replace('\xad', '').replace('\xa0', ' ')).strip()
