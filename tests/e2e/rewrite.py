"""Live «someone else's text» check (task 3, §2.5.2): an LLM rewrites an anonymized document
into a memo keeping the placeholders — the way an external LLM behaves — and the app
restores it. Run by scripts/e2e_private.py --rewrite when Ollama is available.

Checks: no known token left; every name put into new text is a form of a person of the
session; every company put into new text is in quotes; no unknown tokens except ones the
LLM invented (counted apart).
"""
import json
import re
from pathlib import Path

from docx import Document

PROMPT = ('Перепиши документ ниже в служебную записку руководителю: кратко изложи суть, '
          'перечисли риски и вопросы. Плейсхолдеры в квадратных скобках вида [FIO_1], [YUL_2], '
          '[INN_3] сохраняй точно как есть, ставь их в нужные места предложений, не склоняй и '
          'не раскрывай. Ответ — только текст записки.\n\n')


def ollama_writer(timeout=240):
    from core import llm
    schema = {'type': 'object', 'properties': {'text': {'type': 'string'}}, 'required': ['text']}

    def write(text: str) -> str:
        raw = llm._ollama_generate(PROMPT + text, 'Ты юрист, пишешь по-русски.', timeout=timeout,
                                   schema=schema, num_predict=1500)
        return json.loads(raw).get('text', '')
    return write


def run_rewrite(client, src: Path, work: Path, write) -> dict:
    from core.anonymizer import TOKEN_LOOSE_RE, _token_key
    from core.cases import same_person_form
    from tests.e2e.independent import places

    def read_text(p):
        return '\n'.join(places(p)['body'])
    work.mkdir(parents=True, exist_ok=True)
    rep = {'errors': [], 'names': 0, 'orgs': 0, 'invented': 0}
    sid = client.session('rewrite')
    r = client.process(sid, src, spacy=True)
    anon = client.download(sid, r['output'], work / ('anon' + Path(r['output']).suffix))
    text = read_text(anon)[:6000]
    memo = write(text)
    if not memo.strip():
        rep['errors'].append('llm_empty')
        return rep
    d = Document()
    for line in memo.split('\n'):
        d.add_paragraph(line)
    reply = work / 'memo.docx'
    d.save(reply)
    known = {m['token'] for m in client.mappings(sid, archive=True)}
    rep['invented'] = len({_token_key(m) for m in TOKEN_LOOSE_RE.finditer(memo)} - known)
    rr = client.process(sid, reply, mode='deanonymize', name='memo.docx')
    back = client.download(sid, rr['output'], work / 'memo_restored.docx')
    restored = read_text(back)
    left = {_token_key(m) for m in TOKEN_LOOSE_RE.finditer(restored)} & known
    if left:
        rep['errors'].append(f'tokens_left:{len(left)}')
    pv = client.json('get', f'/api/sessions/{sid}/preview/{rr["output"]}')
    marks = (pv.get('restore') or {}).get('marks') or []
    persons = [m['original_form'] for m in client.mappings(sid, archive=True) if m['entity_type'] == 'ФИО']
    for mark, value in marks:
        if mark == 'case':
            rep['names'] += 1
            if not any(same_person_form(p, value) for p in persons):
                rep['errors'].append('name_not_a_form')
    orgs = {m['original_form'] for m in client.mappings(sid, archive=True) if m['entity_type'] == 'ЮЛ'}
    for mark, value in marks:
        if mark == 'exact' and value not in orgs and any(o in value for o in orgs):
            rep['orgs'] += 1
            if not re.search(r'[«"“]', value):
                rep['errors'].append('org_without_quotes')
    return rep
