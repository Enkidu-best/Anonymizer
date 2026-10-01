"""Restore by place, not by occurrence number (task 2, §1.3).

* a deleted or edited mapping must not break restoring a file downloaded earlier;
* paragraphs reordered / removed by an external LLM keep their exact forms;
* changed paragraphs get the main form and are counted as «check the case»;
* tokens unknown to the session are reported, not skipped silently.
"""
from docx import Document

from core.db import delete_mapping, get_session_mappings, get_or_create_token
from core.handlers import process_uploaded_file

PARAS = [
    'Продавец: Белозёров Аркадий Львович, ИНН 711417075366.',
    'Покупатель: ООО «Уральская Логистика» в лице Скворцовой Елены Дмитриевны.',
    'Белозёрову Аркадию Львовичу передано уведомление.',
    'Скворцова Е.Д. подписала акт.',
    'Подпись: А.Л. Белозёров',
]


def _docx(path, paras):
    d = Document()
    for p in paras:
        d.add_paragraph(p)
    d.save(path)


def _texts(path):
    return [p.text for p in Document(path).paragraphs]


def _anon(tmp_path, db, sid):
    src = tmp_path / 'contract.docx'
    _docx(src, PARAS)
    (tmp_path / 'a').mkdir(exist_ok=True)
    r = process_uploaded_file(src, tmp_path / 'a', sid, db, 'anonymize', use_spacy=False)
    return tmp_path / 'a' / r['output_filename']


def _deanon(tmp_path, db, sid, path, name=None):
    (tmp_path / 'b').mkdir(exist_ok=True)
    if name:
        renamed = path.with_name(name)
        renamed.write_bytes(path.read_bytes())
        path = renamed
    r = process_uploaded_file(path, tmp_path / 'b', sid, db, 'deanonymize')
    return r, _texts(tmp_path / 'b' / r['output_filename'])


def test_restore_after_mapping_deleted(tmp_path, tmp_db, session_id):
    anon = _anon(tmp_path, tmp_db, session_id)
    inn = next(m for m in get_session_mappings(tmp_db, session_id) if m['entity_type'] == 'ИНН')
    delete_mapping(tmp_db, session_id, inn['token'])
    assert not any(m['token'] == inn['token'] for m in get_session_mappings(tmp_db, session_id))
    _, back = _deanon(tmp_path, tmp_db, session_id, anon)
    assert back == PARAS


def test_restore_after_mapping_edited(tmp_path, tmp_db, session_id):
    anon = _anon(tmp_path, tmp_db, session_id)
    fio = next(m for m in get_session_mappings(tmp_db, session_id) if m['entity_type'] == 'ФИО')
    # what the PATCH route does: retire the old token, create a new one
    from core.db import retire_mapping
    retire_mapping(tmp_db, session_id, fio['token'])
    get_or_create_token(tmp_db, session_id, 'Иной Человек', 'Иной Человек', 'ФИО')
    _, back = _deanon(tmp_path, tmp_db, session_id, anon)
    assert back == PARAS


def test_llm_reordered_paragraphs_keep_exact_forms(tmp_path, tmp_db, session_id):
    anon = _anon(tmp_path, tmp_db, session_id)
    paras = _texts(anon)
    reply = tmp_path / 'reply.docx'
    # an LLM answer: paragraphs reversed, one dropped, one new sentence with a token
    tok = paras[0].split('[')[1].split(']')[0]
    _docx(reply, list(reversed(paras[1:])) + [f'Риск: [{tok}] может не передать долю.'])
    r, back = _deanon(tmp_path, tmp_db, session_id, reply)
    assert back[:4] == list(reversed(PARAS[1:]))          # exact forms: «Белозёрову Аркадию…», «А.Л. …»
    assert back[4] == 'Риск: Белозёров Аркадий Львович может не передать долю.'
    assert r['restore']['check_case'] >= 1                 # new sentence: main form, flagged
    assert not any('[' in p for p in back)


def test_unknown_tokens_reported(tmp_path, tmp_db, session_id):
    _anon(tmp_path, tmp_db, session_id)
    reply = tmp_path / 'reply.docx'
    _docx(reply, ['Сторона [FIO_99] и [YUL_42] согласовали.'])
    r, back = _deanon(tmp_path, tmp_db, session_id, reply)
    assert set(r['restore']['unknown']) == {'FIO_99', 'YUL_42'}
