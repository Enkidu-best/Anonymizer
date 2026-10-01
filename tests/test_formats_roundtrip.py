"""File-level round trip: anonymize → nothing leaks → deanonymize → identical text."""
from docx import Document
from openpyxl import Workbook, load_workbook

from core.handlers import process_uploaded_file
from tests.test_holdout_contract import CONTRACT, MUST_MASK

INN = '1869879737'


def _run(src, tmp_path, tmp_db, session_id, mode):
    out_dir = tmp_path / mode
    out_dir.mkdir(exist_ok=True)
    r = process_uploaded_file(src, out_dir, session_id, tmp_db, mode, use_spacy=False)
    return out_dir / r['output_filename']


def test_xlsx_numeric_inn_is_masked_and_restored(tmp_path, tmp_db, session_id):
    src = tmp_path / 'in.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.append(['ИНН', int(INN)])
    ws.append(['Контакт', 'Белозёров Аркадий Львович'])
    wb.save(src)

    anon = _run(src, tmp_path, tmp_db, session_id, 'anonymize')
    vals = [c.value for row in load_workbook(anon).active.iter_rows() for c in row]
    assert int(INN) not in vals and INN not in map(str, vals)
    assert 'Белозёров Аркадий Львович' not in vals

    back = _run(anon, tmp_path, tmp_db, session_id, 'deanonymize')
    vals = [c.value for row in load_workbook(back).active.iter_rows() for c in row]
    assert vals == ['ИНН', int(INN), 'Контакт', 'Белозёров Аркадий Львович']


def test_docx_contract_round_trip(tmp_path, tmp_db, session_id):
    src = tmp_path / 'contract.docx'
    doc = Document()
    for line in CONTRACT.split('\n'):
        doc.add_paragraph(line)
    doc.save(src)

    anon = _run(src, tmp_path, tmp_db, session_id, 'anonymize')
    text = '\n'.join(p.text for p in Document(anon).paragraphs)
    assert not [m for m in MUST_MASK if m in text]

    back = _run(anon, tmp_path, tmp_db, session_id, 'deanonymize')
    assert '\n'.join(p.text for p in Document(back).paragraphs) == CONTRACT


def test_rtf_cyrillic_round_trip_with_uc0(tmp_path, tmp_db, session_id):
    """RTF with \\uc0 (no fallback chars) and \\'hh cp1251 escapes: masked and restored."""
    from core.rtf import decode
    name = 'Белозёров Аркадий Львович'
    u = ''.join(f'\\u{ord(c)}' + (' ' if ord(c) > 127 else '') if ord(c) > 127 else c for c in name)
    cp = ''.join(f"\\'{b:02x}" for b in 'ИНН 1869879737'.encode('cp1251'))
    raw = ('{\\rtf1\\ansi\\ansicpg1251\\uc0 {\\info{\\author Иванов}}'
           f'Продавец {u}, {cp}\\par}}').encode('latin-1', errors='ignore')
    src = tmp_path / 'doc.rtf'
    src.write_bytes(raw)
    (tmp_path / 'a').mkdir()
    (tmp_path / 'b').mkdir()
    r = process_uploaded_file(src, tmp_path / 'a', session_id, tmp_db, 'anonymize', use_spacy=False)
    anon_text, _ = decode((tmp_path / 'a' / r['output_filename']).read_bytes())
    assert 'Белозёров' not in anon_text and '1869879737' not in anon_text
    r2 = process_uploaded_file(tmp_path / 'a' / r['output_filename'], tmp_path / 'b', session_id, tmp_db, 'deanonymize')
    back, _ = decode((tmp_path / 'b' / r2['output_filename']).read_bytes())
    assert back == decode(raw)[0]


def test_broken_text_layer_detected():
    from core.handlers import text_layer_ok
    garbage = ('AoroBopy apeHAbl Hex[noro noMerqeH],re or <01> oxrs6pn 2013 roAa r. Mocxaa '
               'O6qecrBo c orpaHhqenxofr orBercrBeHHocrbro nMeHyeMoe B AaflbHefrurevr Bacrnueaofr Enenur '
               'MrxafrnoBHbr gefrcreypulero Ha ocHoBaHLAtA ycraBa')
    good = ('Общество с ограниченной ответственностью, именуемое в дальнейшем Арендодатель, в лице '
            'генерального директора, действующего на основании устава, с одной стороны, и Арендатор, '
            'заключили настоящее соглашение о нижеследующем')
    assert not text_layer_ok(garbage)
    assert text_layer_ok(good)
    assert text_layer_ok('Share purchase agreement between the Seller and the Buyer dated twelve March '
                         'two thousand and twenty five regarding the shares of the Company and other terms')


def test_textutil_markup_normalized_and_props_scrubbed():
    from lxml import etree
    from core.ooxml import normalize_wordml, W
    xml = (f'<w:document xmlns:w="{W}"><w:body><w:p><w:pPr><w:jc w:val="both"/><w:ind w:left="1"/></w:pPr>'
           '<w:r><w:rPr><w:rFonts w:ascii="X"/><w:sz w:val="24"/><w:sz-cs w:val="24"/><w:b/><w:color w:val="0"/>'
           '</w:rPr><w:t>x</w:t></w:r></w:p></w:body></w:document>')
    root = etree.fromstring(xml)
    normalize_wordml(root)
    rpr = [c.tag.split('}')[1] for c in root.iter(f'{{{W}}}rPr').__next__()]
    ppr = [c.tag.split('}')[1] for c in root.iter(f'{{{W}}}pPr').__next__()]
    assert rpr == ['rFonts', 'b', 'color', 'sz', 'szCs']
    assert ppr == ['ind', 'jc']


def test_docx_title_and_subject_cleared(tmp_path, tmp_db, session_id):
    from docx import Document
    src = tmp_path / 'p.docx'
    d = Document()
    d.add_paragraph('Текст договора')
    d.core_properties.title = 'Тверская обл., г. Примерск, пр. Ленина, д. 12'
    d.core_properties.subject = 'Тел./факс: (48242)3-33-08'
    d.save(src)
    (tmp_path / 'a').mkdir()
    r = process_uploaded_file(src, tmp_path / 'a', session_id, tmp_db, 'anonymize', use_spacy=False)
    props = Document(tmp_path / 'a' / r['output_filename']).core_properties
    assert not props.title and not props.subject
