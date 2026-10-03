"""Restoring NEW text of an external LLM (task 3, §2.2, §2.3). Fictional data.

companies: in quotes → the name; without quotes → «name»; standing alone → «ООО «name»»;
persons: case from the word on the left; marked 'case' and highlighted yellow in Word;
numbers, addresses, companies are never «check the case».
"""
from docx import Document

from core.anonymizer import anonymize_text, make_rev_finder, apply_spans
from core.handlers import process_uploaded_file

SRC = ('Продавец: ООО «Вектор Трейд», ИНН 7707083893, в лице Белозёрова Аркадия Львовича. '
       'Адрес: г. Москва, ул. Лесная, д. 5.')


def _setup(db, sid):
    out, occ = anonymize_text(SRC, db, sid, use_spacy=False)
    toks = {t.split('_')[0]: t for t in occ}
    return out, toks


def _restore(db, sid, text, settings=None):
    find = make_rev_finder(db, sid, settings=settings)
    find.prepare(text)
    return apply_spans(text, find(text)), find


def test_org_quotes_and_legal_form(tmp_db, session_id):
    _, t = _setup(tmp_db, session_id)
    y = t['YUL']
    cases = {
        f'Договор с ООО «[{y}]» расторгнут.': 'Договор с ООО «Вектор Трейд» расторгнут.',
        f'Компания [{y}] не ответила.': 'Компания «Вектор Трейд» не ответила.',
        f'ООО [{y}] не ответило.': 'ООО «Вектор Трейд» не ответило.',
        f'Риск: [{y}] может не оплатить.': 'Риск: ООО «Вектор Трейд» может не оплатить.',
        f'Требование к **[{y}]** не предъявлено.': 'Требование к **ООО «Вектор Трейд»** не предъявлено.',
        f'Требование к "[{y}]" не предъявлено.': 'Требование к "Вектор Трейд" не предъявлено.',
    }
    for src, want in cases.items():
        got, find = _restore(tmp_db, session_id, src)
        assert got == want, (src, got)
        assert find.stats['check_case'] == 0              # a company is never «check the case»
    got, _ = _restore(tmp_db, session_id, f'Риск: [{y}] может.', {'org_quotes': 'keep', 'org_opf': 'never'})
    assert got == 'Риск: Вектор Трейд может.'


def test_person_case_in_new_text(tmp_db, session_id):
    _, t = _setup(tmp_db, session_id)
    f = t['FIO']
    cases = {
        f'Письмо направить [{f}] до пятницы.': 'Письмо направить Белозёрову Аркадию Львовичу до пятницы.',
        f'Подлинник хранится у [{f}].': 'Подлинник хранится у Белозёрова Аркадия Львовича.',
        f'Договор заключён с {f}.': 'Договор заключён с Белозёровым Аркадием Львовичем.',
        f'[{f}] подписал акт.': 'Белозёров Аркадий Львович подписал акт.',
        f'Сведения о [{f.replace("_", " ")}] устарели.': 'Сведения о Белозёрове Аркадии Львовиче устарели.',
    }
    for src, want in cases.items():
        got, find = _restore(tmp_db, session_id, src)
        assert got == want, (src, got)
        assert find.stats['check_case'] == 1 and find.marks[0][0] == 'case'


def test_numbers_and_addresses_not_flagged(tmp_db, session_id):
    _, t = _setup(tmp_db, session_id)
    got, find = _restore(tmp_db, session_id, f'Проверить ИНН [{t["INN"]}] и адрес [{t["ADR"]}].')
    assert '7707083893' in got and 'Лесная' in got
    assert find.stats['check_case'] == 0
    assert [m for m, _ in find.marks] == ['exact', 'exact']


def test_llm_case_answer_checked(tmp_db, session_id, monkeypatch):
    """Rules can't tell → one LLM request; an answer that is not a form of the name is refused."""
    import core.llm as llm
    _, t = _setup(tmp_db, session_id)
    f = t['FIO']
    calls = []

    def fake(items, timeout=40):
        calls.append(items)
        return {0: 'Белозёрова Аркадия Львовича', 1: 'Скворцовой Елене'}
    monkeypatch.setattr(llm, 'choose_cases', fake)
    text = f'Риск связан с тем, что согласие [{f}] не получено.\nИ ещё вопрос: что если [{f}] возразит, то как?'
    # the second place: «что если [FIO] возразит» — the rules decide nominative? no — let the LLM answer
    got, find = _restore(tmp_db, session_id, text)
    assert len(calls) <= 1
    assert 'согласие Белозёрова Аркадия Львовича' in got
    assert 'Скворцовой' not in got


def test_word_highlight_for_case_places(tmp_path, tmp_db, session_id):
    _, t = _setup(tmp_db, session_id)
    d = Document()
    p = d.add_paragraph()
    p.add_run(f'Письмо направить [{t["FIO"]}] и копию в [{t["YUL"]}].')
    src = tmp_path / 'reply.docx'
    d.save(src)
    (tmp_path / 'o').mkdir()
    r = process_uploaded_file(src, tmp_path / 'o', session_id, tmp_db, 'deanonymize')
    back = Document(tmp_path / 'o' / r['output_filename'])
    runs = back.paragraphs[0].runs
    assert back.paragraphs[0].text == 'Письмо направить Белозёрову Аркадию Львовичу и копию в ООО «Вектор Трейд».'
    hl = [x.text for x in runs if x.font.highlight_color is not None]
    assert hl == ['Белозёрову Аркадию Львовичу']
    assert r['restore']['check_case'] == 1
