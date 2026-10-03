"""«Точно»: the file is ready after the rules, the LLM reviews in the background and the
user applies its proposal (task 2, §3.2). Ollama is mocked."""
import importlib
import io
import json
import os
import sys
import time

from docx import Document


def test_background_review_and_apply(tmp_path, monkeypatch):
    os.environ['ANONYMIZER_DATA_DIR'] = str(tmp_path / 'data')
    sys.modules.pop('app', None)
    app_mod = importlib.import_module('app')
    import core.llm as llm

    def fake(prompt, system, timeout=90, schema=None, num_ctx=None, num_predict=2048):
        if schema and 'drop' in schema.get('properties', {}):
            # short answer: only changes by number — item «Арендатор» is a role, a name was missed
            found = prompt.split('СТРОКИ:')[0].split('\n')
            drop = [int(l.split('.')[0]) for l in found if 'Арендатор»' in l and l[:1].isdigit()]
            add = [{'text': 'Агафонов', 'type': 'FIO'}] if 'Агафонов' in prompt.split('СТРОКИ:')[1] else []
            return json.dumps({'drop': drop, 'merge': [], 'add': add})
        return json.dumps({'entities': []})
    monkeypatch.setattr(llm, '_ollama_generate', fake)
    monkeypatch.setattr(llm, 'check_ollama', lambda: {'available': True})
    import core.anonymizer as A
    monkeypatch.setattr(A, '_ner_ready', False)   # spaCy (if loaded by other tests) would find the name itself

    c = app_mod.app.test_client()
    sid = c.post('/api/sessions', json={'name': 'x'}).get_json()['id']
    # a role wrongly recorded as a person — the review must propose to unmask it
    from core.db import get_or_create_token
    get_or_create_token(app_mod.DB_PATH, sid, 'Арендатор', 'Арендатор', 'ФИО')
    src = tmp_path / 'd.docx'
    d = Document()
    d.add_paragraph('Арендатор Тарханов Глеб Игоревич обязуется оплатить.')
    d.add_paragraph('Свидетель: эксперт Агафонов, «Независимая оценка» Иркутск.')
    d.save(src)
    r = c.post('/api/process', data={'mode': 'anonymize', 'session_id': sid, 'engine': 'accurate',
                                     'files': (io.BytesIO(src.read_bytes()), 'd.docx')},
               content_type='multipart/form-data').get_json()['results'][0]
    assert r['status'] == 'ok' and r['llm_job']
    for _ in range(100):
        st = c.get(f'/api/llm-jobs/{r["llm_job"]}').get_json()
        if st['status'] != 'running':
            break
        time.sleep(0.05)
    assert st['status'] == 'done', st
    prop = st['proposal']
    assert [x['value'] for x in prop['add']] == ['Агафонов']
    assert [x['value'] for x in prop['remove']] == ['Арендатор']
    res = c.post(f'/api/llm-jobs/{r["llm_job"]}/apply',
                 json={'add': [0], 'remove': [x['token'] for x in prop['remove']]}).get_json()
    assert res['ok']
    out = c.get(f'/api/download/{sid}/{res["output"]}')
    text = '\n'.join(p.text for p in Document(io.BytesIO(out.data)).paragraphs)
    assert 'Агафонов' not in text and 'Тарханов' not in text
    assert text.startswith('Арендатор ')
