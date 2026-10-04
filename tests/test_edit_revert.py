"""Editing a record and then editing it back (owner, v3.4.0): the value must stay masked,
also after processing the file again. Fictional data."""
import importlib
import io
import os
import sys

from docx import Document


def _client(tmp_path):
    os.environ['ANONYMIZER_DATA_DIR'] = str(tmp_path / 'data')
    sys.modules.pop('app', None)
    app_mod = importlib.import_module('app')
    import core.anonymizer as A
    A._ner_ready = False
    c = app_mod.app.test_client()
    sid = c.post('/api/sessions', json={'name': 'x'}).get_json()['id']
    return c, sid


def _process(c, sid, path):
    data = {'mode': 'anonymize', 'session_id': sid, 'engine': 'fast', 'use_spacy': 'false',
            'files': (io.BytesIO(path.read_bytes()), path.name)}
    r = c.post('/api/process', data=data, content_type='multipart/form-data').get_json()
    return r['results'][0]


def test_edit_then_edit_back(tmp_path):
    c, sid = _client(tmp_path)
    src = tmp_path / 'd.docx'
    d = Document()
    d.add_paragraph('Адрес объекта: г. Тверь, ул. Озёрная, д. 17. Продавец Белозёров Аркадий Львович.')
    d.save(src)
    _process(c, sid, src)
    url = f'/api/sessions/{sid}/mappings'
    m = next(x for x in c.get(url).get_json() if x['entity_type'] == 'ФИО')
    orig = m['original_form']
    r1 = c.patch(f'{url}/{m["token"]}', json={'original_form': orig + 'а', 'on_duplicate': 'separate'}).get_json()
    assert r1['ok']
    r2 = c.patch(f'{url}/{r1["new_token"]}', json={'original_form': orig, 'on_duplicate': 'merge'}).get_json()
    assert r2['ok']
    active = {x['original_form'] for x in c.get(url).get_json()}
    assert orig in active, active                     # was: silently gone
    r = _process(c, sid, src)                         # process again: still masked
    assert r['status'] == 'ok'
    out = tmp_path / 'out.docx'
    out.write_bytes(c.get(f'/api/download/{sid}/{r["output"]}').data)
    text = '\n'.join(p.text for p in Document(out).paragraphs)
    assert 'Белозёров' not in text, text
