"""Сценарии окна Anonymizer 4.0 - номера из «Функциональное описание.md» v2.1 (UC-K…, UC-A…, UC-M…, UC-U1, UC-P…,
UC-D…, UC-R1, UC-S…, UC-W…, UC-N…, UC-G1), раздел 8.

Свой сервер на порту 8074 (стандарт 6.3) и своя временная папка данных (6.10); серверы с подменой (модель имён
загружается долго или не загружается, медленная обработка, подменная Ollama) - на 8075-8079. Только вымышленные
документы. Chromium; основные сценарии - ещё и WebKit. Браузер и окна Андрея не открываются (правило 10).

    .venv312/bin/python -m pytest -m ui tests/test_ui_v4.py
"""
import http.server
import json
import os
import re
import sqlite3
import subprocess
import sys
import threading
import time
import urllib.parse
import urllib.request
from pathlib import Path

import pytest
from docx import Document

pytestmark = pytest.mark.ui
try:
    from playwright.sync_api import sync_playwright
except ImportError:
    pytest.skip('Playwright не установлен', allow_module_level=True)

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tests' / 'bench'))
import ids  # noqa: E402

ids.R.seed(4002)
INN = ids.inn10()
CONTRACT = ['ДОГОВОР ПОСТАВКИ № 17', 'г. Москва, 3 сентября 2026 г.',
            f'ООО «Северный ветер», ИНН {INN}, именуемое «Поставщик», в лице генерального директора Петрова Ивана '
            'Сергеевича, действующего на основании Устава, и Сидорова Анна Викторовна, именуемая «Покупатель», '
            'заключили договор.',
            'Поставщик передаёт товар. Приёмку ведёт кладовщик Ковалёв Дмитрий Олегович.',
            'Подписи: Петров И. С. / Сидорова А. В.']
PORT = 8074


# ── вымышленные документы ────────────────────────────────────────────────────
def make_docx(path: Path, paragraphs=CONTRACT) -> Path:
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path


def make_txt(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding='utf-8')
    return path


def make_xlsx(path: Path, rows=1000) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(['№', 'Плательщик', 'Сумма'])
    for i in range(rows):
        ws.append([i + 1, 'Петров Иван Сергеевич' if i % 10 == 0 else 'ООО «Северный ветер»', 1000 + i])
    wb.save(str(path))
    return path


# ── сервер и браузер ─────────────────────────────────────────────────────────
def _get(url, timeout=1):
    with urllib.request.urlopen(url, timeout=timeout) as r:
        return json.loads(r.read())


def _call(base, path, data=None, method=None):
    req = urllib.request.Request(base + path, data=json.dumps(data).encode() if data is not None else None,
                                 method=method or ('POST' if data is not None else 'GET'),
                                 headers={'Content-Type': 'application/json'})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read() or b'{}')


class Server:
    def __init__(self, work: Path, port: int, env=None):
        self.work, self.port, self.url = work, port, f'http://127.0.0.1:{port}'
        self.data = work / 'data'
        self.env = dict(os.environ, ANONYMIZER_DATA_DIR=str(self.data), ANONYMIZER_LOG_DIR=str(work / 'logs'),
                        ANONYMIZER_PORT=str(port), ANONYMIZER_NO_OPEN='1', HEAVY_LOCK_WHO='background',
                        **(env or {}))
        self.proc = None

    def start(self):
        self.log = open(self.work / f'server-{self.port}.log', 'a')
        self.proc = subprocess.Popen([sys.executable, str(ROOT / 'app.py')], cwd=str(ROOT), env=self.env,
                                     stdout=self.log, stderr=subprocess.STDOUT)
        end = time.time() + 60
        while time.time() < end:
            try:
                _get(self.url + '/api/info')
                return self
            except Exception:
                time.sleep(0.25)
        raise RuntimeError('сервер не запустился: ' + (self.work / f'server-{self.port}.log').read_text()[-2000:])

    def stop(self):
        if self.proc:
            self.proc.terminate()
            try:
                self.proc.wait(10)
            except subprocess.TimeoutExpired:
                self.proc.kill()
            self.log.close()

    def wait_ner(self, timeout=90):
        end = time.time() + timeout
        while time.time() < end:
            st = _get(self.url + '/api/status')
            if st['ner']['ready'] or st['ner']['error']:
                return st['ner']
            time.sleep(0.3)

    def api(self, path, data=None, method=None):
        return _call(self.url, path, data, method)


@pytest.fixture(scope='module')
def server(tmp_path_factory):
    s = Server(tmp_path_factory.mktemp('ui'), PORT).start()
    s.wait_ner()
    yield s
    s.stop()


@pytest.fixture(scope='module')
def pw():
    with sync_playwright() as p:
        yield p


@pytest.fixture(scope='module')
def chromium(pw):
    b = pw.chromium.launch()
    yield b
    b.close()


@pytest.fixture(scope='module')
def webkit(pw):
    try:
        b = pw.webkit.launch()
    except Exception as ex:
        pytest.skip(f'WebKit недоступен: {ex}')
    yield b
    b.close()


def open_page(browser, url, width=1440, height=900, clipboard=False):
    ctx = browser.new_context(viewport={'width': width, 'height': height}, locale='ru-RU', accept_downloads=True,
                              permissions=['clipboard-read', 'clipboard-write'] if clipboard else [])
    pg = ctx.new_page()
    pg.set_default_timeout(30000)
    pg.errors = []
    pg.on('pageerror', lambda e: pg.errors.append(str(e)))
    pg.goto(url)
    pg.wait_for_function("document.getElementById('model-t').textContent !== '…'")
    return pg


@pytest.fixture
def page(chromium, server):
    pg = open_page(chromium, server.url)
    yield pg
    assert not pg.errors, pg.errors
    pg.context.close()


@pytest.fixture(params=['chromium', 'webkit'])
def any_page(request, server):
    br = request.getfixturevalue(request.param)
    pg = open_page(br, server.url)
    yield pg
    assert not pg.errors, pg.errors
    pg.context.close()


def upload(pg, *paths):
    pg.set_input_files('#file-input', [str(p) for p in paths])
    pg.wait_for_function(f"document.querySelectorAll('.row[data-input]').length + document.querySelectorAll('[data-rej]').length >= {len(paths)}")


def process(pg):
    pg.click('#go-anon')
    pg.wait_for_selector('#sum-anon', timeout=120000)


def anon_one(pg, tmp_path, name='Договор поставки - Петров.docx', paragraphs=CONTRACT):
    upload(pg, make_docx(tmp_path / name, paragraphs))
    process(pg)
    return pg.evaluate('S.sid')


def files_of(server, sid):
    return server.api(f'/api/sessions/{sid}/files')['files']


def mappings(server, sid, all_=False):
    return server.api(f'/api/sessions/{sid}/mappings' + ('?all=1' if all_ else ''))


def docx_text(server, sid, name) -> str:
    import io
    data = urllib.request.urlopen(f'{server.url}/api/download/{sid}/{urllib.parse.quote(name)}').read()
    return '\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def toast_text(pg, part=''):
    pg.wait_for_selector(f'.toast:has-text("{part}")' if part else '.toast')
    return pg.inner_text('.toast')


def crumbs(pg):
    return pg.evaluate("[...document.getElementById('crumbs').children].map(e => e.textContent.trim()).join(' ')")


# ── 3.1 Каркас ───────────────────────────────────────────────────────────────
def test_K1_start_empty_version_text_size(any_page, server):
    """UC-K1: запуск - «Обезличить», пусто, без сессии; версия внизу меню = /api/info; основной текст 17 px."""
    pg = any_page
    assert pg.inner_text('#content h1') == 'Обезличить'
    assert 'Сессия появится сама при первом файле' in pg.inner_text('#content')
    assert pg.evaluate('S.sid') is None
    v = server.api('/api/info')['version']
    assert pg.inner_text('#ver') == 'Anonymizer ' + re.sub(r'\.0$', '', v)
    assert pg.evaluate("getComputedStyle(document.body).fontSize") == '17px'
    assert crumbs(pg) == 'Обезличить'
    assert pg.is_disabled('#go-anon') and 'Добавьте файлы' in pg.inner_text('.head .act')


def test_K2_theme_cycle(page):
    """UC-K2: «Тема» → светлая → тёмная → системная; выбор запоминается."""
    seq = []
    for _ in range(3):
        page.click('#mi-theme')
        seq.append(page.evaluate("document.documentElement.dataset.theme || 'system'"))
    assert seq == ['light', 'dark', 'system']
    page.click('#mi-theme')
    page.reload()
    page.wait_for_selector('#content h1')
    assert page.evaluate("document.documentElement.dataset.theme") == 'light'


def test_K3_english_window(page):
    """UC-K3 (кратко; полностью - scripts/ui_check.py): EN - в окне нет кириллицы, RU - снова русский."""
    page.click('#lang [data-lang=en]')
    page.wait_for_function("document.querySelector('#content h1').textContent === 'Anonymize'")
    for view in ('anon', 'sessions', 'dict', 'settings', 'help'):
        page.evaluate(f"go('{view}')")
        page.wait_for_timeout(300)
        txt = page.inner_text('body')
        assert not re.search('[А-Яа-яЁё]', txt), (view, re.findall('[А-Яа-яЁё]+', txt)[:5])
    page.click('#lang [data-lang=ru]')
    page.wait_for_function("document.querySelector('#content h1').textContent === 'Справка'")


def test_K4_narrow_window(chromium, server):
    """UC-K4: 900 px - меню значками, путь и шапка не обрезаны."""
    pg = open_page(chromium, server.url, width=900)
    assert pg.evaluate("document.getElementById('side').getBoundingClientRect().width") == 60
    assert pg.evaluate("[...document.querySelectorAll('.mi span')].every(s => getComputedStyle(s).display === 'none')")
    cut = pg.evaluate("""[...document.querySelectorAll('body *')].filter(e => e.children.length === 0 && e.textContent.trim()
        && getComputedStyle(e).overflow !== 'visible' && e.scrollWidth > e.clientWidth + 1).map(e => e.textContent.trim())""")
    assert not cut, cut
    pg.context.close()


def test_K5_zoom_persists(page):
    """UC-K5: «Aa» → 150 % → после перезапуска окна 150 %; Cmd+0 → 100 %."""
    page.click('#zoom')
    for _ in range(3):                 # 100 → 110 → 125 → 150 %
        page.click('#z-inc')
    assert page.inner_text('#zoom-val') == '150 %'
    page.reload()
    page.wait_for_selector('#content h1')
    assert page.evaluate("getComputedStyle(document.body).fontSize") == '25.5px'
    page.keyboard.press('Meta+0')
    assert page.evaluate("getComputedStyle(document.body).fontSize") == '17px'


def test_version_mismatch_banner(chromium, server):
    """Стандарт 6.5: версия сервера ≠ версии окна - «Программа обновилась - перезапустите» с кнопкой."""
    ctx = chromium.new_context(locale='ru-RU')
    pg = ctx.new_page()
    pg.route('**/api/info', lambda r: r.fulfill(json={'name': 'Anonymizer', 'version': '9.9.9', 'migration': {}}))
    pg.goto(server.url)
    pg.wait_for_selector('#banner:not(.hidden)')
    assert 'Программа обновилась - перезапустите' in pg.inner_text('#banner')
    assert pg.is_visible('#b-restart')
    ctx.close()


# ── 3.2 Обезличить ───────────────────────────────────────────────────────────
def test_A1_docx_name_without_surname(any_page, server, tmp_path):
    """UC-A1 + правка 6: «Готово: 1 файл · утечек нет»; в окне - исходное имя, в имени файла - только метки;
    сессия «Сессия ДД.ММ» в «Сессиях»."""
    pg = any_page
    sid = anon_one(pg, tmp_path)
    assert re.search(r'Готово: 1 файл · \d+ запис[ьией]+ · утечек нет', pg.inner_text('#sum-anon'))
    row = pg.inner_text('.row[data-input]')
    assert 'Договор поставки - Петров.docx' in row
    out = files_of(server, sid)[0]['output']
    assert 'Петров' not in out and '[FIO_' in out and out in row
    assert re.fullmatch(r'Обезличить › Сессия \d\d\.\d\d(?: \(\d+\))? · 1 файл', crumbs(pg))
    pg.click('[data-go=sessions]')
    pg.wait_for_selector('#s-list .row')
    assert re.search(r'Сессия \d\d\.\d\d', pg.inner_text('#s-list'))


def test_A2_progress_stop_continue(chromium, tmp_path):
    """UC-A2: 3 файла → полоса с процентом и этапом → «Остановить» → «Остановлено: готово N из 3» → остаток."""
    s = Server(tmp_path, 8075, {'ANONYMIZER_SLOW': '1.5'}).start()
    try:
        s.wait_ner()
        pg = open_page(chromium, s.url)
        upload(pg, *[make_txt(tmp_path / f'письмо {i}.txt', f'Письмо {i}. Ответственный - Петров Иван Сергеевич.') for i in range(3)])
        pg.click('#go-anon')
        pg.wait_for_selector('#prog .bar')
        assert re.search(r'\d+ %', pg.inner_text('#prog')) and 'Файл 1 из 3' in pg.inner_text('#prog')
        assert 'Работает: файл' in pg.inner_text('#model')
        pg.click('#p-stop')
        pg.wait_for_selector('text=Остановлено: готово', timeout=30000)
        m = re.search(r'готово (\d) из 3', pg.inner_text('#content'))
        done = int(m.group(1))
        assert 1 <= done < 3
        assert pg.inner_text('#go-anon') == f'Обезличить ({3 - done})'
        pg.click('#go-anon')
        pg.wait_for_function("S.files.filter(f => f.status === 'done').length === 3", timeout=60000)
        assert not pg.errors
        pg.context.close()
    finally:
        s.stop()


def test_A3_waits_for_name_model(chromium, tmp_path):
    """UC-A3: модель распознавания имён ещё грузится → «Готовлю распознавание имён…» → обработка началась сама."""
    s = Server(tmp_path, 8076, {'ANONYMIZER_NER_DELAY': '6'}).start()
    try:
        pg = open_page(chromium, s.url)
        upload(pg, make_docx(tmp_path / 'Договор.docx'))
        assert 'Готовлю распознавание имён' in pg.inner_text('#model')
        pg.click('#go-anon')
        pg.wait_for_selector('text=Готовлю распознавание имён… обработка начнётся сама')
        pg.wait_for_selector('#sum-anon', timeout=120000)
        f = files_of(s, pg.evaluate('S.sid'))[0]
        assert f['status'] == 'done' and not f['result']['without_ner']
        pg.context.close()
    finally:
        s.stop()


def test_A4_name_model_failed(chromium, tmp_path):
    """UC-A4: модель не загрузилась → «Повторить» и красная «Обезличить без распознавания имён» с предупреждением."""
    s = Server(tmp_path, 8077, {'ANONYMIZER_NER_FAIL': '1'}).start()
    try:
        s.wait_ner()
        pg = open_page(chromium, s.url)
        assert 'Ошибка модели' in pg.inner_text('#model')
        upload(pg, make_docx(tmp_path / 'Договор.docx'))
        pg.click('#go-anon')
        pg.wait_for_selector('#p-noner')
        txt = pg.inner_text('#prog')
        assert 'Модель распознавания имён не загрузилась' in txt and 'проверьте каждый файл' in txt
        assert pg.is_visible('#p-retry')
        pg.click('#p-noner')
        pg.wait_for_selector('#sum-anon', timeout=60000)
        assert 'без распознавания имён' in pg.inner_text('.row[data-input]')
        pg.context.close()
    finally:
        s.stop()


def test_A5_reopen_sections_after_restart(chromium, tmp_path):
    """UC-A5: обработать, вернуть данные, перезапустить, открыть сессию → файлы каждого раздела на месте с итогами."""
    s = Server(tmp_path, 8078).start()
    try:
        s.wait_ner()
        pg = open_page(chromium, s.url)
        sid = anon_one(pg, tmp_path)
        out = files_of(s, sid)[0]['output']
        reply = tmp_path / 'reply' / out
        reply.parent.mkdir()
        reply.write_bytes(urllib.request.urlopen(f'{s.url}/api/download/{sid}/{urllib.parse.quote(out)}').read())
        pg.click('[data-go=restore]')
        pg.click('[data-mode=file]')
        upload(pg, reply)
        pg.click('#go-restore')
        pg.wait_for_selector('#files-restore [data-act=preview]', timeout=60000)
        pg.context.close()
        s.stop()
        s.start()
        s.wait_ner()
        pg = open_page(chromium, s.url)
        pg.click('[data-go=sessions]')
        pg.click('#s-list .row')
        pg.wait_for_selector('#sum-anon')
        rows = pg.locator('#files-anon .row').all_inner_texts()
        assert len(rows) == 1 and 'Договор поставки - Петров.docx' in rows[0] and 'Найдено' in rows[0]
        pg.click('[data-go=restore]')
        pg.click('[data-mode=file]')
        pg.wait_for_selector('#files-restore .row')
        rows = pg.locator('#files-restore .row').all_inner_texts()
        assert len(rows) == 1 and 'данные возвращены' in rows[0]
        pg.context.close()
    finally:
        s.stop()


def test_A6_too_big_file(page, tmp_path):
    """UC-A6: файл 120 МБ → «Файл больше 100 МБ - разделите его», остальные обработаны."""
    big = tmp_path / 'Большой.docx'
    with open(big, 'wb') as f:
        f.truncate(120 * 1024 * 1024)
    upload(page, big, make_docx(tmp_path / 'Обычный.docx'))
    assert 'Файл больше 100 МБ - разделите его' in page.inner_text('#files-anon')
    process(page)
    assert 'Найдено' in page.inner_text('.row[data-input="Обычный.docx"]')


def test_A7_remove_from_session(page, server, tmp_path):
    """UC-A7: «Убрать из сессии» → окно программы → файла нет в списке и на диске, записи таблицы замен на месте."""
    sid = anon_one(page, tmp_path)
    before = len(mappings(server, sid))
    page.click('.row[data-input] [data-act=more]')
    page.click('.pop .mitem.danger')
    page.wait_for_selector('.dlg')
    assert 'Записи таблицы замен останутся' in page.inner_text('.dlg')
    page.click('.dlg .btn.primary')
    page.wait_for_function("!document.querySelector('.row[data-input]')")
    assert files_of(server, sid) == []
    assert not list((server.data / 'uploads' / sid / 'input').iterdir())
    assert len(mappings(server, sid)) == before


def test_A8_copy_instruction(chromium, server, tmp_path):
    """UC-A8 + правка 3: «Скопировать инструкцию для нейросети» → в буфере текст инструкции."""
    pg = open_page(chromium, server.url, clipboard=True)
    anon_one(pg, tmp_path)
    assert pg.inner_text('#instr') == 'Скопировать инструкцию для нейросети'
    pg.click('#instr')
    assert 'Скопировано - вставьте перед документом' in toast_text(pg, 'Скопировано')
    clip = pg.evaluate('navigator.clipboard.readText()')
    assert '[FIO_1]' in clip and 'Не изменяй' in clip
    pg.context.close()


# ── 3.3 Таблица замен, 3.4 окно записи, 3.5 «Не маскировать» ──────────────────
def test_M1_duplicate_value_dialog(page, server, tmp_path):
    """UC-M1: «+ Добавить вручную» значение, которое уже есть другой формой → «Такое значение уже есть» с выбором."""
    anon_one(page, tmp_path)
    page.click('#tbl-add')
    page.fill('#rec-v', 'Петрову И. С.')
    page.keyboard.press('Enter')
    page.wait_for_selector('text=Такое значение уже есть')
    btns = page.locator('.dlg .btns button').all_inner_texts()
    assert btns == ['Отмена', 'Всё равно добавить отдельно', 'Объединить с ним']
    page.click('.dlg .btn.primary')
    assert 'Добавлено' in toast_text(page, 'Добавлено')


def test_M2_add_before_first_file(page, server, tmp_path):
    """UC-M2: «+ Добавить вручную» до первого файла → запись в таблице; после обработки значение - меткой."""
    page.click('#tbl-add')
    assert 'Добавится и применится при обработке' in page.inner_text('.dlg')
    page.fill('#rec-v', 'Северный причал')
    page.select_option('#rec-t', 'ЮЛ')
    page.click('.dlg .btn.primary')
    page.wait_for_selector('#tbl-body tr[data-token]')
    assert 'Северный причал' in page.inner_text('#tbl-body')
    sid = anon_one(page, tmp_path, 'Склад.docx', CONTRACT + ['Склад - Северный причал.'])
    out = [f for f in files_of(server, sid) if f['input'] == 'Склад.docx'][0]['output']
    assert 'Северный причал' not in docx_text(server, sid, out)


def test_M3_unmask_filter_mask_again(page, server, tmp_path):
    """UC-M3: «Не маскировать» → «Не маскируется» → «Вернуть маску» → в файле снова метка, та же."""
    sid = anon_one(page, tmp_path)
    tok = next(m['token'] for m in mappings(server, sid) if m['original_form'].startswith('Ковалёв'))
    page.click(f'#tbl-body tr[data-token="{tok}"] [data-row-more]')
    page.click('.pop .mitem >> text=Не маскировать')
    page.wait_for_selector('.dlg')
    assert page.locator('.dlg input[value=session]').is_checked()
    page.keyboard.press('Enter')
    assert 'больше не маскируется' in toast_text(page, 'больше не маскируется')
    page.wait_for_function("!S.rebuild")
    out = files_of(server, sid)[0]['output']
    assert 'Ковалёв Дмитрий Олегович' in docx_text(server, sid, out)
    page.select_option('#tbl-type', '__excluded')
    assert tok in page.inner_text('#tbl-body')
    page.click(f'#tbl-body tr[data-token="{tok}"] [data-row-more]')
    page.click('.pop .mitem >> text=Вернуть маску')
    page.wait_for_function(f"S.maps.some(m => m.token === '{tok}') && !S.rebuild")
    out = files_of(server, sid)[0]['output']
    assert f'[{tok}]' in docx_text(server, sid, out)


def test_M4_type_filter_all_30(page, tmp_path):
    """UC-M4: «Вид ▾» - 30 видов и «Не маскируется»; ИНН находит свои строки."""
    anon_one(page, tmp_path)
    opts = page.locator('#tbl-type option').all_inner_texts()
    assert len(opts) == 32 and opts[0] == 'Вид: все' and opts[-1] == 'Не маскируется'
    page.select_option('#tbl-type', 'ИНН')
    rows = page.locator('#tbl-body tr').all_inner_texts()
    assert rows and all('ИНН' in r for r in rows)


def test_column_initial_form(page, tmp_path):
    """Правка 2: колонка «Начальная форма» с подсказкой; ФИО - в именительном падеже."""
    anon_one(page, tmp_path)
    th = page.locator('.tbl th').nth(2)
    assert th.inner_text().startswith('Начальная форма')
    assert th.get_attribute('title') == 'ФИО - в именительном падеже, программа склонит сама'
    assert 'Петров Иван Сергеевич' in page.inner_text('#tbl-body')


def test_dialog_keys(page):
    """4.4: окна программы - Esc и щелчок мимо закрывают, не `confirm()`."""
    page.click('#tbl-add')
    page.keyboard.press('Escape')
    assert not page.is_visible('.dlg')
    page.click('#tbl-add')
    page.mouse.click(5, 5)
    assert not page.is_visible('.dlg')
    page.click('#tbl-add')
    page.keyboard.press('Enter')
    assert page.inner_text('.dlg .err') == 'Заполните значение'
    page.keyboard.press('Escape')


def test_U1_unmask_always_in_preview(page, server, tmp_path):
    """UC-U1: в просмотре «Не маскировать» с «всегда» → Словарь «Никогда не маскировать» → новая сессия не маскирует."""
    anon_one(page, tmp_path)
    page.click('.row[data-input] [data-act=preview]')
    page.wait_for_selector('#pv-list .val')
    page.click('#pv-list .val:has-text("Ковалёв")')
    page.click('#pv-list [data-un]')
    page.check('.dlg input[value=always]')
    page.click('.dlg .btn.primary')
    assert 'больше не маскируется' in toast_text(page, 'больше не маскируется')
    never = server.api('/api/dictionary')['never']
    assert [x['value'] for x in never] == ['Ковалёв Дмитрий Олегович']
    page.click('.toast button')
    page.wait_for_function("!S.rebuild && S.maps.some(m => m.original_form.startsWith('Ковалёв'))")
    assert server.api('/api/dictionary')['never'] == []


# ── 3.6 Просмотр ─────────────────────────────────────────────────────────────
def test_P1_preview_navigation(any_page, server, tmp_path):
    """UC-P1: документ с чипами, справа группы с числами; щелчок → к месту, ↓ → следующее; «← Назад» → список."""
    pg = any_page
    anon_one(pg, tmp_path)
    pg.click('.row[data-input] [data-act=preview]')
    pg.wait_for_selector('#pv-body .tok')
    assert 'Обезличить › Сессия' in crumbs(pg) and 'Договор поставки - Петров.docx' in crumbs(pg)
    assert 'Скачается как:' in pg.inner_text('.pvbar')
    assert 'ЛЮДИ' in pg.inner_text('#pv-list').upper()
    pg.click('#pv-list .val:has-text("Петров")')
    first = pg.evaluate("[...document.querySelectorAll('#pv-body .tok')].findIndex(e => e.classList.contains('cur'))")
    pg.keyboard.press('ArrowDown')
    second = pg.evaluate("[...document.querySelectorAll('#pv-body .tok')].findIndex(e => e.classList.contains('cur'))")
    assert first >= 0 and second > first
    pg.click('#back')
    pg.wait_for_selector('#files-anon .row')


def test_P2_mask_selection(page, server, tmp_path):
    """UC-P2: выделить слово → «Замаскировать» → «Сохранить» → на месте слова метка."""
    sid = anon_one(page, tmp_path)
    page.click('.row[data-input] [data-act=preview]')
    page.wait_for_selector('#pv-body .tok')
    page.evaluate("""() => { const w = document.createTreeWalker(document.getElementById('pv-body'), NodeFilter.SHOW_TEXT);
        while (w.nextNode()) { const n = w.currentNode, i = n.nodeValue.indexOf('Устава');
          if (i >= 0) { const r = document.createRange(); r.setStart(n, i); r.setEnd(n, i + 6);
            const s = getSelection(); s.removeAllRanges(); s.addRange(r); break; } } }""")
    page.dispatch_event('#pv-body', 'mouseup')
    page.click('.floatbtn')
    assert page.input_value('#rec-v') == 'Устава'
    page.select_option('#rec-t', 'ЮЛ')
    page.click('.dlg .btn.primary')
    page.wait_for_function("S.maps.some(m => m.original_form === 'Устава') && !S.rebuild")
    out = files_of(server, sid)[0]['output']
    assert 'Устава' not in docx_text(server, sid, out)


def test_P3_duplicates_distinct(page, server, tmp_path):
    """UC-P3: две похожие организации → «Проверка дублей» → «Это разные» → пара больше не предлагается."""
    sid = anon_one(page, tmp_path, 'Две компании.docx', CONTRACT + ['Агент - ООО «Вектор Трейд».', 'Перевозчик - ООО «Вектор Тренд».'])
    page.click('.row[data-input] [data-act=preview]')
    page.wait_for_selector('details.dups')
    page.click('details.dups summary')
    page.click('[data-distinct="0"]')
    assert 'это разные' in toast_text(page, 'это разные')
    assert server.api(f'/api/sessions/{sid}/duplicates') == []


def test_P4_excel_truncated_note(page, tmp_path):
    """UC-P4: Excel 1 000 строк → «Показаны первые 800 строк…»."""
    upload(page, make_xlsx(tmp_path / 'Реестр.xlsx'))
    process(page)
    page.click('.row[data-input] [data-act=preview]')
    page.wait_for_selector('.xl-cut')
    assert 'Показаны первые 800 строк и 40 столбцов листа' in page.inner_text('.xl-cut')


# ── 3.7-3.8 Вернуть данные ───────────────────────────────────────────────────
def test_D4_paste_text_restore_and_copy(chromium, server, tmp_path):
    """UC-D4, UC-D2 (вопрос 1): вставить ответ → падеж подобран (жёлтым), чужая метка (красным) → «Скопировать»."""
    pg = open_page(chromium, server.url, clipboard=True)
    sid = anon_one(pg, tmp_path)
    fio = next(m['token'] for m in mappings(server, sid) if m['original_form'].startswith('Петрова Ивана'))
    pg.click('[data-go=restore]')
    pg.click('[data-mode=text]')
    pg.fill('#rs-text', f'Направляем [{fio}] ответ на претензию.\nКопия - [FIO_99].')
    pg.click('#go-restore')
    pg.wait_for_selector('#rs-result')
    assert pg.inner_text('#rs-result .rv-case') == 'Петрову Ивану Сергеевичу'
    assert pg.inner_text('#rs-result .rv-unknown') == '[FIO_99]'
    assert 'Не найдены в этой сессии: [FIO_99]' in pg.inner_text('#content')
    pg.click('#rs-copy')
    clip = pg.evaluate('navigator.clipboard.readText()')
    assert clip.startswith('Направляем Петрову Ивану Сергеевичу ответ')
    pg.click('[data-go=anon]')
    pg.click('[data-go=restore]')
    assert pg.input_value('#rs-text') == ''           # текст не хранится
    assert not pg.errors
    pg.context.close()


def test_D3_two_sessions_choice(page, server, tmp_path):
    """UC-D3: две сессии → «Сессия ▾» → вторая → значения из второй."""
    anon_one(page, tmp_path)
    page.click('#sess-new')
    anon_one(page, tmp_path, 'Второй.docx', ['Договор подписал Лебедев Пётр Андреевич.'])
    s1, s2 = server.api('/api/sessions')[1], server.api('/api/sessions')[0]
    page.click('[data-go=restore]')
    page.click('[data-mode=text]')
    page.select_option('#rs-sess', s2['id'])
    page.fill('#rs-text', 'Подписал [FIO_1].')
    page.click('#go-restore')
    page.wait_for_selector('#rs-result')
    assert 'Лебедев' in page.inner_text('#rs-result')
    page.select_option('#rs-sess', s1['id'])
    page.fill('#rs-text', 'Подписал [FIO_1].')
    page.click('#go-restore')
    page.wait_for_function("document.getElementById('rs-result') && !document.getElementById('rs-result').textContent.includes('Лебедев')")


def test_D5_R1_restore_file_name_and_colours(page, server, tmp_path):
    """UC-D5 (правка 1), UC-R1: «… (данные возвращены).docx»; в просмотре - точно зелёным, падеж жёлтым."""
    sid = anon_one(page, tmp_path)
    fio = next(m['token'] for m in mappings(server, sid) if m['original_form'].startswith('Петрова Ивана'))
    reply = make_docx(tmp_path / 'reply' / 'Ответ нейросети - претензия.docx', [f'Направляем [{fio}] ответ.', 'Покупатель - [FIO_2].'])
    page.click('[data-go=restore]')
    page.click('[data-mode=file]')
    upload(page, reply)
    page.click('#go-restore')
    page.wait_for_selector('#files-restore [data-act=preview]', timeout=60000)
    out = [f for f in files_of(server, sid) if f['section'] == 'restore'][0]['output']
    assert out == 'Ответ нейросети - претензия (данные возвращены).docx'
    assert re.search(r'Вернулось \d+: \d+ точно, \d+ с подбором падежа', page.inner_text('#files-restore'))
    page.click('#files-restore [data-act=preview]')
    page.wait_for_selector('#pv-body .rv-case')
    assert 'Петрову Ивану Сергеевичу' in page.inner_text('#pv-body .rv-case')
    assert 'Вернулось' in page.inner_text('.pvbar')
    assert not page.is_visible('#pv-orig')


def test_no_sessions_message(chromium, tmp_path):
    """3.7: сессий нет - «Сначала обезличьте файлы…»."""
    s = Server(tmp_path, 8079).start()
    try:
        pg = open_page(chromium, s.url)
        pg.click('[data-go=restore]')
        pg.wait_for_selector('text=Сначала обезличьте файлы')
        pg.click('[data-go=sessions]')
        pg.wait_for_selector('text=Сессий нет')
        pg.click('[data-go=dict]')
        pg.wait_for_selector('text=Пусто. Сюда попадают значения')
        pg.context.close()
    finally:
        s.stop()


# ── 3.9-3.10 Сессии ──────────────────────────────────────────────────────────
def test_S1_S2_S4_sessions_list(page, server, tmp_path):
    """UC-S1, UC-S2, UC-S4: «Сессия ДД.ММ» с «1 файл · N записей»; «⋯ → Переименовать»; поиск."""
    sid = anon_one(page, tmp_path)
    page.click('[data-go=sessions]')
    row = page.locator(f'#s-list .row[data-id="{sid}"]')
    assert re.search(r'1 файл · \d+ запис', row.inner_text())
    row.locator('[data-more]').click()
    page.click('.pop .mitem >> text=Переименовать')
    page.fill('#s-list input.field', 'Аренда офиса')
    page.keyboard.press('Enter')
    page.wait_for_selector('#s-list .nm >> text=Аренда офиса')
    page.fill('#s-q', 'аренд')
    assert page.locator('#s-list .row').count() == 1
    page.click('#s-list .row')
    page.wait_for_selector('#sum-anon')
    assert 'Аренда офиса' in crumbs(page)


def test_S3_retention_mark_and_delete(page, server, tmp_path):
    """UC-S3: старше 27 дней - «удалится через 3 дня»; старше срока - сессии, папки и записей нет."""
    sid = anon_one(page, tmp_path)
    con = sqlite3.connect(str(server.data / 'anon.db'))
    con.execute("UPDATE sessions SET created_at = datetime('now','localtime','-27 days','-1 hours') WHERE id=?", (sid,))
    con.commit()
    page.click('[data-go=sessions]')
    assert 'удалится через 3 дня' in page.inner_text(f'#s-list .row[data-id="{sid}"]')
    con.execute("UPDATE sessions SET created_at = datetime('now','localtime','-31 days') WHERE id=?", (sid,))
    con.commit()
    server.api('/api/settings', {'retention_days': 90})
    server.api('/api/settings', {'retention_days': 30})         # смена срока - удаление по сроку сразу
    assert sid not in {x['id'] for x in server.api('/api/sessions')}
    assert not (server.data / 'uploads' / sid).exists()
    assert con.execute('SELECT COUNT(*) FROM mappings WHERE session_id=?', (sid,)).fetchone()[0] == 0
    con.close()


def test_S5_delete_dialog(page, server, tmp_path):
    """UC-S5: ✕ → окно с числами → Esc и Enter не удаляют; «Удалить навсегда» → нет в списке, папки и записей."""
    sid = anon_one(page, tmp_path)
    page.click('[data-go=sessions]')
    page.click(f'#s-list .row[data-id="{sid}"] [data-del]')
    txt = page.inner_text('.dlg')
    assert re.search(r'1 файл, \d+ запис', txt) and 'Отменить удаление нельзя' in txt and 'Словарь остаётся' in txt
    assert page.evaluate("document.activeElement.textContent") == 'Отмена'
    page.keyboard.press('Escape')
    page.click(f'#s-list .row[data-id="{sid}"] [data-del]')
    page.keyboard.press('Enter')
    assert sid in {x['id'] for x in server.api('/api/sessions')}
    page.click(f'#s-list .row[data-id="{sid}"] [data-del]')
    page.click('.dlg .btn.danger')
    assert 'Сессия удалена' in toast_text(page, 'Сессия удалена')
    assert sid not in {x['id'] for x in server.api('/api/sessions')}
    assert not (server.data / 'uploads' / sid).exists()
    con = sqlite3.connect(str(server.data / 'anon.db'))
    assert con.execute('SELECT COUNT(*) FROM mappings WHERE session_id=?', (sid,)).fetchone()[0] == 0
    con.close()


# ── 3.11 Словарь ─────────────────────────────────────────────────────────────
def test_W1_W2_dictionary(page, server, tmp_path):
    """UC-W1: «всегда» при добавлении → в «Всегда маскировать» → маскируется в новой сессии; UC-W2: ✕ и «Вернуть»."""
    page.click('#tbl-add')
    page.fill('#rec-v', 'Зелёная долина')
    page.select_option('#rec-t', 'ЮЛ')
    page.check('.dlg input[value=always]')
    page.click('.dlg .btn.primary')
    page.click('#sess-new')
    sid = anon_one(page, tmp_path, 'Новый.docx', CONTRACT + ['Субподрядчик - Зелёная долина.'])
    assert 'Зелёная долина' not in docx_text(server, sid, files_of(server, sid)[0]['output'])
    page.click('[data-go=dict]')
    assert 'Зелёная долина' in page.inner_text('#content .list')
    assert 'добавлено' in page.inner_text('#content .list') and 'вручную' in page.inner_text('#content .list')
    page.click('#content .list [data-rm]')
    assert 'убран из Словаря' in toast_text(page, 'убран из Словаря')
    page.click('.toast button')
    page.wait_for_selector('#content .list .nm >> text=Зелёная долина')
    page.click('#content .list [data-rm]')
    page.wait_for_timeout(5500)
    assert not page.is_visible('.toast')
    assert server.api('/api/dictionary')['always'] == []


# ── 3.12 Настройки ───────────────────────────────────────────────────────────
class FakeOllama(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        body = json.dumps({'models': [{'name': 'qwen3.5:9b'}, {'name': 'llama3:8b'}]}).encode()
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.end_headers()
        self.wfile.write(json.dumps({'response': '{"entities": []}', 'done': True}).encode())

    def log_message(self, *a):
        pass


@pytest.fixture
def fake_ollama():
    srv = http.server.ThreadingHTTPServer(('127.0.0.1', 0), FakeOllama)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f'http://127.0.0.1:{srv.server_address[1]}'
    srv.shutdown()


def test_N1_settings_saved_at_once(page):
    """UC-N1: «Хранить сессии» → 7 → после перезапуска окна 7; «Сохранено» у параметра; правка 4 - подпись."""
    page.click('#mi-settings')
    page.click('[data-ret="7"]')
    page.wait_for_selector('#saved-retention_days.show')
    page.reload()
    page.click('#mi-settings')
    assert 'on' in page.get_attribute('[data-ret="7"]', 'class')
    page.click('[data-grp=more]')
    d = page.inner_text('.srow:has-text("Использовать нейросеть")')
    assert 'Проверяет найденное и подбирает падеж в спорных местах. Медленнее, работает в очереди видеокарты' in d
    assert 'пользы' not in d
    page.click('[data-ret="30"]')


def test_N2_model_choice_kept(chromium, tmp_path, fake_ollama):
    """UC-N2: «Использовать нейросеть» вкл → появилась «Модель»; модель выбрана → после перезапуска та же."""
    s = Server(tmp_path, 8075, {'ANONYMIZER_OLLAMA_URL': fake_ollama}).start()
    try:
        pg = open_page(chromium, s.url)
        pg.click('#mi-settings')
        pg.click('[data-grp=more]')
        assert not pg.is_visible('[data-sel=llm_model]')
        pg.click('[data-sw=use_llm]')
        pg.wait_for_selector('[data-sel=llm_model]')
        pg.select_option('[data-sel=llm_model]', 'llama3:8b')
        pg.wait_for_selector('#saved-llm_model.show')
        pg.context.close()
        s.stop()
        s.start()
        assert _get(s.url + '/api/settings')['llm_model'] == 'llama3:8b'
        pg = open_page(chromium, s.url)
        pg.click('#mi-settings')
        pg.click('[data-grp=more]')
        pg.wait_for_selector('[data-sel=llm_model]')
        assert pg.input_value('[data-sel=llm_model]') == 'llama3:8b'
        s.api('/api/settings', {'use_llm': False})
        pg.context.close()
    finally:
        s.stop()


def test_N3_selects_not_cut(page, tmp_path):
    """UC-N3, стандарт 4.4: выпадающие списки не обрезают выбранный текст."""
    anon_one(page, tmp_path)
    for view in ('anon', 'settings'):
        page.evaluate(f"go('{view}')")
        page.wait_for_timeout(300)
        cut = page.evaluate("[...document.querySelectorAll('select')].filter(s => s.scrollWidth > s.clientWidth + 1).map(s => s.id || s.dataset.sel)")
        assert not cut, (view, cut)


# ── 3.14 Видеокарта занята, 3.15 предложения нейросети ───────────────────────
def test_G1_gpu_busy_window_and_cancel_wait(chromium, tmp_path, fake_ollama):
    """UC-G1: нейросеть включена, замок у «Herald» → окно с 4 кнопками → «Подождать очереди» → «ждёт видеокарту…
    Отменить ожидание» → проверки нет, файл после правил готов."""
    lock = tmp_path / 'lock' / '.heavy-lock'
    lock.parent.mkdir()
    s = Server(tmp_path, 8076, {'ANONYMIZER_OLLAMA_URL': fake_ollama, 'ANONYMIZER_HEAVY_LOCK': str(lock)}).start()
    holder = ROOT / 'common' / 'lock' / 'tests' / '_holder.py'
    hd = subprocess.Popen([sys.executable, str(holder), str(lock), 'Herald', 'human', 'hold', 'озвучка книги'],
                          stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)
    try:
        assert hd.stdout.readline().strip() == 'taken'
        s.wait_ner()
        s.api('/api/settings', {'use_llm': True})
        pg = open_page(chromium, s.url)
        pg.wait_for_function("document.getElementById('model-t').textContent.includes('Herald')")
        assert pg.inner_text('#model-t') == 'Нейросеть: занята «Herald» - в очереди'
        upload(pg, make_docx(tmp_path / 'Договор.docx'))
        pg.click('#go-anon')
        pg.wait_for_selector('.hl-dialog')
        txt = pg.inner_text('.hl-dialog')
        assert 'Видеокарта занята: «Herald»' in txt and 'озвучка книги' in txt
        assert pg.eval_on_selector_all('.hl-dialog button', 'bs => bs.map(b => b.dataset.answer)') == ['wait', 'now', 'cpu', 'cancel']
        pg.click('.hl-dialog button[data-answer=wait]')
        pg.wait_for_selector('#sum-anon', timeout=60000)
        pg.wait_for_selector('text=Нейросеть ждёт видеокарту: «Herald» - начнётся сама', timeout=20000)
        pg.click('[data-llm-cancel]')
        pg.wait_for_selector('text=Проверка нейросетью отменена - файл после правил готов', timeout=20000)
        assert files_of(s, pg.evaluate('S.sid'))[0]['status'] == 'done'
        pg.context.close()
    finally:
        hd.stdin.write('release\n')
        hd.stdin.flush()
        hd.wait(10)
        s.stop()


def test_llm_proposals_unmask_unchecked(page, tmp_path):
    """3.15 (вопрос 5, было Р-25): «Снять маску» по умолчанию не отмечено, «Замаскировать» и «Объединить» - отмечены."""
    anon_one(page, tmp_path)
    page.evaluate("""S.llm['x'] = {status: 'done', proposal: {add: [{value: 'Котов', type: 'ФИО'}],
        remove: [{token: 'FIO_1', value: 'Петров', category: 'ROLE'}], merge: [{token: 'FIO_2', into: 'FIO_1', value: 'a', into_value: 'b'}]}};
        setTimeout(() => openReview('x'))""")
    page.wait_for_selector('.dlg')
    assert page.inner_text('.dlg h2') == 'Предложения нейросети'
    assert page.is_checked('.dlg input[data-kind=add]')
    assert not page.is_checked('.dlg input[data-kind=remove]')
    assert page.is_checked('.dlg input[data-kind=merge]')
    page.keyboard.press('Escape')
