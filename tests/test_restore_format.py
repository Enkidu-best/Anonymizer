"""Restored values keep the formatting of the token's run (task 3, §2.1). Fictional data.

A value is written into the run that held the token (its w:rPr: size, font, bold…), in
body text of 8 / 11 / 14 pt, a table, a header and a heading; a token cut into runs by an
external tool («[» in a tiny run, «FIO_1]» in the normal one) takes the run with most of it.
"""
from docx import Document
from docx.shared import Pt

from core.db import get_or_create_token
from core.handlers import process_uploaded_file

NAME = 'Белозёров Аркадий Львович'
ORG = 'Вектор Трейд'


def _session(db, sid):
    assert get_or_create_token(db, sid, NAME, NAME, 'ФИО') == 'FIO_1'
    assert get_or_create_token(db, sid, ORG, ORG, 'ЮЛ') == 'YUL_1'


def _runs_with(doc_or_part, needle):
    out = []
    paras = list(doc_or_part.paragraphs)
    for t in getattr(doc_or_part, 'tables', []):
        for row in t.rows:
            for c in row.cells:
                paras += c.paragraphs
    for p in paras:
        for r in p.runs:
            if needle in r.text:
                out.append((r, p))
    return out


def _size(run, para):
    if run.font.size:
        return run.font.size.pt
    return para.style.font.size.pt if para.style.font.size else None


def test_value_keeps_run_size_font_bold(tmp_path, tmp_db, session_id):
    _session(tmp_db, session_id)
    d = Document()
    d.sections[0].header.paragraphs[0].add_run('Исполнитель: [FIO_1]').font.size = Pt(8)
    d.add_heading('Договор с [YUL_1]', level=1)
    for pt in (8, 11, 14):
        p = d.add_paragraph()
        p.add_run('Сторона: ').font.size = Pt(10)
        r = p.add_run(f'[FIO_1]')
        r.font.size = Pt(pt)
        r.font.name = 'Arial'
        r.bold = pt == 14
        p.add_run(' и дальше.').font.size = Pt(10)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).paragraphs[0].add_run('Сторона').font.size = Pt(9)
    t.cell(0, 1).paragraphs[0].add_run('[YUL_1]').font.size = Pt(9)
    # an external tool cut the token: «[» tiny, the rest normal
    p = d.add_paragraph()
    p.add_run('Подписал ').font.size = Pt(12)
    p.add_run('[').font.size = Pt(4)
    p.add_run('FIO_1]').font.size = Pt(12)
    src = tmp_path / 'reply.docx'
    d.save(src)
    (tmp_path / 'out').mkdir()
    r = process_uploaded_file(src, tmp_path / 'out', session_id, tmp_db, 'deanonymize')
    back = Document(tmp_path / 'out' / r['output_filename'])

    body = _runs_with(back, NAME)
    sizes = [_size(r, p) for r, p in body]
    assert sizes == [8, 11, 14, 12], sizes
    assert [r.font.name for r, _ in body[:3]] == ['Arial'] * 3
    assert [bool(r.bold) for r, _ in body[:3]] == [False, False, True]
    head = _runs_with(back.sections[0].header, NAME)
    assert [_size(r, p) for r, p in head] == [8]
    cell = _runs_with(back, ORG)
    assert any(_size(r, p) == 9 for r, p in cell)
    heading = [p for p in back.paragraphs if p.style.name.startswith('Heading') and ORG in p.text]
    assert heading
    assert '[' not in ''.join(p.text for p in back.paragraphs)
