"""
Anonymizer 4.0 - сервер на 127.0.0.1:8070; окно - Chrome в режиме приложения (desktop/launcher.sh, стандарт 6.2).

    python app.py             сервер (порт 8070; ANONYMIZER_PORT - другой, тесты - 8074)
    ./start.sh                то же и окно; NO_BROWSER=1 - без окна (проверки)

Данные - ~/Library/Application Support/Anonymizer (platformdirs, и из исходников тоже; ANONYMIZER_DATA_DIR - своя
папка для тестов), журнал - ~/Library/Logs/Anonymizer (стандарт 6.4).
"""

import io
import os
import shutil
import sys
import threading
import time
import uuid
import zipfile
from pathlib import Path

from flask import Flask, request, jsonify, send_file, send_from_directory

from core.version import __version__ as APP_VERSION
from core import session_files as sf

BUNDLE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BUNDLE_DIR / 'static'
PORT = int(os.environ.get('ANONYMIZER_PORT') or 8070)
MAX_MB = 100


def _user_data_dir() -> Path:
    """Per-user writable folder, never next to the program (standard 6.4, was Р-35)."""
    try:
        from platformdirs import user_data_dir
        return Path(user_data_dir('Anonymizer', appauthor=False))
    except ImportError:
        if sys.platform == 'darwin':
            return Path.home() / 'Library' / 'Application Support' / 'Anonymizer'
        if sys.platform == 'win32':
            return Path(os.environ.get('APPDATA', Path.home())) / 'Anonymizer'
        return Path.home() / '.local' / 'share' / 'Anonymizer'


DATA_DIR = Path(os.environ.get('ANONYMIZER_DATA_DIR') or _user_data_dir())
UPLOADS_DIR = DATA_DIR / 'uploads'
DB_PATH = DATA_DIR / 'anon.db'
DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOADS_DIR.mkdir(exist_ok=True)

# ── Journal (no personal data) ───────────────────────────────────────────────
from core import log as journal  # noqa: E402
journal.setup(journal.default_dir())
from core import settings as user_settings  # noqa: E402
user_settings.configure(DATA_DIR)

# ── Database: the move to 4.0 with a backup and a check of the numbers (description 5.4) ──
from core.db import init_db  # noqa: E402
from core import migrate  # noqa: E402
MIGRATION = migrate.migrate(DATA_DIR, init_db)
if MIGRATION.get('done'):
    user_settings.update({'dict_migrated': int(MIGRATION.get('moved') or 0)})
    journal.event('migrated_4_0', moved=MIGRATION.get('moved'), counts=MIGRATION.get('after'))
elif MIGRATION.get('error'):
    journal.event('migration_failed', error=MIGRATION['error'][:200])


# ── Retention: a session is deleted whole after the term (description 3.9, was Р-38) ──
def purge_expired() -> int:
    from core.db import old_sessions, delete_session, get_all_sessions
    days = user_settings.get()['retention_days']
    gone = 0
    for sid in old_sessions(DB_PATH, days):
        delete_session(DB_PATH, sid)
        shutil.rmtree(UPLOADS_DIR / sid, ignore_errors=True)
        gone += 1
    known = {s['id'] for s in get_all_sessions(DB_PATH)}
    limit = time.time() - days * 86400
    for d in UPLOADS_DIR.iterdir():          # folders of sessions that are no longer in the base
        try:
            if d.is_dir() and d.name not in known and d.stat().st_mtime < limit:
                shutil.rmtree(d, ignore_errors=True)
        except OSError:
            pass
    if gone:
        journal.event('sessions_expired', count=gone, days=days)
    return gone


purge_expired()


def _purge_daily():
    while True:
        time.sleep(24 * 3600)
        try:
            purge_expired()
        except Exception as ex:
            journal.error('purge_failed', ex)


threading.Thread(target=_purge_daily, daemon=True).start()


def _startup_journal():
    info = {}
    if user_settings.get()['use_llm']:
        try:
            from core.llm import check_ollama
            st = check_ollama()
            info = {'ollama': st.get('available'), 'models': st.get('models')}
        except Exception:
            pass
    journal.startup_info(APP_VERSION, info)


threading.Thread(target=_startup_journal, daemon=True).start()

# ── The name model loads in the background; the processing waits for it (description 3.2) ──
from core.anonymizer import start_ner_loading  # noqa: E402
start_ner_loading()


def _llm_setup():
    """The neural network only when «Использовать нейросеть» is on (description 3.12): the model from the settings,
    warm-up only with a free GPU (core/llm.preload_model_async)."""
    st = user_settings.get()
    from core import llm
    if st['llm_model']:
        llm.set_model(st['llm_model'])
    if st['use_llm']:
        llm.preload_model_async()


_llm_setup()

# ── Flask app ─────────────────────────────────────────────────────────────────
app = Flask(__name__, static_folder=str(STATIC_DIR))
app.config['MAX_CONTENT_LENGTH'] = (MAX_MB + 1) * 1024 * 1024


@app.before_request
def _local_only():
    """Only the app's own page may talk to the server: a web page opened in any browser could otherwise POST to
    127.0.0.1 (delete sessions, read mappings). Host check also defeats DNS rebinding (description 4.10)."""
    host = (request.host or '').split(':')[0]
    if host not in ('127.0.0.1', 'localhost'):
        return jsonify({'error': 'forbidden host'}), 403
    origin = request.headers.get('Origin')
    if origin and origin != 'null':
        from urllib.parse import urlparse
        o = urlparse(origin)
        if o.hostname not in ('127.0.0.1', 'localhost') or o.port != _port(request.host):
            return jsonify({'error': 'forbidden origin'}), 403
    if request.method not in ('GET', 'HEAD', 'OPTIONS'):
        site = request.headers.get('Sec-Fetch-Site')
        if site and site not in ('same-origin', 'none'):
            return jsonify({'error': 'cross-site request'}), 403


def _port(host: str):
    """Port of «host:port»; None without a port - the page and the server must agree."""
    try:
        return int((host or '').split(':')[1])
    except (IndexError, ValueError):
        return None


@app.errorhandler(413)
def _too_big(_e):
    return jsonify({'error': 'too_big', 'code': 'too_big'}), 413


def _body() -> dict:
    return request.get_json(force=True, silent=True) or {}


def _safe(name: str) -> bool:
    return bool(name) and '..' not in name and '/' not in name and '\\' not in name


# ── Window, version ───────────────────────────────────────────────────────────
@app.route('/')
def index():
    return send_from_directory(STATIC_DIR, 'index.html')


@app.route('/api/info')
def info():
    """Version of the server - the window compares it with its own (standard 6.5, was Р-32)."""
    return jsonify({'name': 'Anonymizer', 'version': APP_VERSION, 'data_dir': str(DATA_DIR),
                    'migration': {k: v for k, v in MIGRATION.items() if k in ('done', 'moved', 'error')}})


@app.route('/api/status')
def status():
    """The line of the model state (standard 2.1): the name model; the neural network - only if switched on."""
    from core.anonymizer import get_ner_status
    st = user_settings.get()
    out = {'version': APP_VERSION, 'ner': get_ner_status(), 'use_llm': st['use_llm']}
    out.update(out['ner'])        # 3.x clients and tests read ready / loading / error at the top
    if st['use_llm']:
        out['llm'] = _llm_state()
        from core import heavy
        h = heavy.status_json('accurate')
        out['gpu'] = {'busy': bool(h.get('holder')) and not h.get('ours'),
                      'program': (h.get('holder') or {}).get('program')}
    return jsonify(out)


_LLM_CACHE = {'t': 0, 'state': None}


def _llm_state() -> dict:
    """Ollama: available and the chosen model present - cached for 20 s (the status is polled)."""
    if time.time() - _LLM_CACHE['t'] > 20 or _LLM_CACHE['state'] is None:
        from core.llm import check_ollama, DEFAULT_MODEL
        s = check_ollama()
        model = user_settings.get()['llm_model'] or DEFAULT_MODEL
        _LLM_CACHE.update(t=time.time(), state={'available': s['available'], 'models': s['models'],
                                                'model': model, 'has_model': model in s['models']})
    return _LLM_CACHE['state']


@app.route('/api/debug')
def debug():
    import platform
    from core.anonymizer import get_ner_status
    try:
        import spacy
        spacy_ver, model_installed = spacy.__version__, spacy.util.is_package('ru_core_news_lg')
    except Exception as e:
        spacy_ver, model_installed = f'IMPORT ERROR: {e}', False
    info_ = {'python_version': sys.version, 'platform': platform.platform(), 'executable': sys.executable,
             'spacy_version': spacy_ver, 'model_installed': model_installed, 'ner_status': get_ner_status(),
             'data_dir': str(DATA_DIR)}
    return '<pre>' + '\n\n'.join(f'{k}: {v}' for k, v in info_.items()) + '</pre>'


@app.route('/api/ner-retry', methods=['POST'])
def ner_retry():
    """«Повторить» when the name model did not load (description 3.2)."""
    from core.anonymizer import retry_ner_loading
    retry_ner_loading()
    return jsonify({'ok': True})


# ── Neural network (Ollama), only with «Использовать нейросеть» ───────────────
@app.route('/api/llm-status')
def llm_status():
    _LLM_CACHE['t'] = 0
    return jsonify(_llm_state())


@app.route('/api/llm-start', methods=['POST'])
def llm_start():
    from core.llm import try_start_ollama
    try_start_ollama()
    _LLM_CACHE['t'] = 0
    return jsonify(_llm_state())


@app.route('/api/llm-progress')
def llm_progress():
    from core.llm import get_llm_progress
    return jsonify(get_llm_progress())


@app.route('/api/llm-cancel', methods=['POST'])
def llm_cancel():
    """«Остановить» and «Отменить ожидание» (description 3.2, 3.14)."""
    from core.llm import cancel_llm
    cancel_llm()
    jid = _body().get('job')
    if jid and jid in REVIEWS and REVIEWS[jid]['status'] in ('running', 'wait_lock'):
        REVIEWS[jid]['cancel'] = True
        CANCEL[jid].set()
    return jsonify({'ok': True})


@app.route('/api/heavy-lock')
def heavy_lock_status():
    """GPU queue (shared module common/lock, standard 2.4): who holds it, the queue; «На процессоре» forecast."""
    from core import heavy
    return jsonify(heavy.status_json(request.args.get('kind', 'accurate')))


@app.route('/api/heavy-lock/page', methods=['POST'])
def heavy_lock_page():
    """«Открыть страницу «Видеокарта»» from the queue window (shared module, README 12)."""
    from common.lock import heavy_lock as hl
    from core import heavy
    return jsonify({'opened': hl.open_page(heavy.path())})


@app.route('/common/lock/<path:name>')
def common_lock_ui(name):
    """The window «Видеокарта занята» - shared module (common/lock/lock_ui)."""
    return send_from_directory(BUNDLE_DIR / 'common' / 'lock' / 'lock_ui', name)


# ── Settings ──────────────────────────────────────────────────────────────────
_USER_KEYS = ('org_quotes', 'org_opf', 'highlight_case', 'pdf_to_word', 'keep_pdf', 'retention_days', 'use_llm',
              'llm_model')


@app.route('/api/settings', methods=['GET', 'POST'])
def settings_route():
    """Every change is saved at once (description 3.12, was Р-28)."""
    if request.method == 'POST':
        data = {k: v for k, v in _body().items() if k in _USER_KEYS}
        before = user_settings.get()
        st = user_settings.update(data)
        if 'llm_model' in data:
            from core.llm import set_model
            set_model(st['llm_model'])
            _LLM_CACHE['t'] = 0
        if st['use_llm'] and not before['use_llm']:
            _LLM_CACHE['t'] = 0
            from core.llm import preload_model_async
            preload_model_async()
        if st['retention_days'] != before['retention_days']:
            purge_expired()
        return jsonify(st)
    return jsonify(user_settings.get())


# ── Sessions ──────────────────────────────────────────────────────────────────
def _session_row(s: dict) -> dict:
    import datetime
    files = sf.load(UPLOADS_DIR / s['id'])
    days = user_settings.get()['retention_days']
    try:
        created = datetime.datetime.strptime(s['created_at'], '%Y-%m-%d %H:%M:%S')
        left = days - (datetime.datetime.now() - created).total_seconds() / 86400
    except (TypeError, ValueError):
        left = days
    return {'id': s['id'], 'name': s['name'], 'created_at': s['created_at'], 'records': s['records'],
            'files': len(files), 'files_anon': sum(f['section'] == 'anon' for f in files),
            'files_restore': sum(f['section'] == 'restore' for f in files),
            'ready_anon': sum(f['section'] == 'anon' and f['status'] == 'done' for f in files),
            'ready_restore': sum(f['section'] == 'restore' and f['status'] == 'done' for f in files),
            'expires_in_days': max(0, int(left + 0.999))}


@app.route('/api/sessions', methods=['GET'])
def list_sessions():
    from core.db import get_all_sessions
    return jsonify([_session_row(s) for s in get_all_sessions(DB_PATH)])


@app.route('/api/sessions', methods=['POST'])
def create_session_route():
    """A session appears by itself with the first file: «Сессия 24.10», never from a file name (was Р-39)."""
    from core.db import create_session, default_session_name
    name = (_body().get('name') or '').strip()[:120] or default_session_name(DB_PATH)
    sid = create_session(DB_PATH, name)
    journal.event('session_create', session=sid[:8])
    return jsonify({'id': sid, 'name': name}), 201


@app.route('/api/sessions/<sid>', methods=['GET'])
def get_session_route(sid):
    from core.db import get_session
    s = get_session(DB_PATH, sid)
    if not s:
        return jsonify({'error': 'not_found'}), 404
    out = _session_row(s)
    out['file_list'] = sf.load(UPLOADS_DIR / sid)
    return jsonify(out)


@app.route('/api/sessions/<sid>', methods=['PATCH'])
def rename_session_route(sid):
    from core.db import rename_session
    name = (_body().get('name') or '').strip()[:120]
    if not name:
        return jsonify({'error': 'empty'}), 400
    rename_session(DB_PATH, sid, name)
    return jsonify({'ok': True, 'name': name})


@app.route('/api/sessions/<sid>', methods=['DELETE'])
def delete_session(sid):
    """At once and for good (description 3.10): originals, files, the replacement table. The Dictionary stays."""
    from core.db import delete_session as db_delete_session
    db_delete_session(DB_PATH, sid)
    journal.event('session_delete', session=sid[:8])
    shutil.rmtree(UPLOADS_DIR / sid, ignore_errors=True)
    return jsonify({'ok': True})


# ── Files of a session ────────────────────────────────────────────────────────
def _dirs(sid, section='anon'):
    """Session folder, uploaded files of the section (originals - input/, neural network replies - replies/), output/."""
    d = UPLOADS_DIR / sid
    return d, d / ('input' if section == 'anon' else 'replies'), d / 'output'


def _unique_input(in_dir: Path, name: str) -> str:
    p = Path(name)
    n = 2
    while (in_dir / name).exists():
        name = f'{p.stem} ({n}){p.suffix}'
        n += 1
    return name


def _store_upload(sid, f, section) -> dict:
    """Saves one uploaded file into input/ and the manifest. Returns the entry or {'code': …}."""
    from core.handlers import safe_upload_name, ALLOWED_EXTENSIONS
    sdir, in_dir, out_dir = _dirs(sid, section)
    in_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    name = safe_upload_name(f.filename or 'document')
    if Path(name).suffix.lower() not in ALLOWED_EXTENSIONS:
        return {'code': 'format', 'input': name}
    data = f.read()
    if len(data) > MAX_MB * 1024 * 1024:
        return {'code': 'too_big', 'input': name}
    files = sf.load(sdir)
    for e in files:
        if e['section'] == section and e['input'] == name and e.get('size') == len(data):
            return {'code': 'duplicate', 'input': name}
    name = _unique_input(in_dir, name)
    try:
        (in_dir / name).write_bytes(data)
    except OSError as ex:
        return {'code': 'disk' if getattr(ex, 'errno', 0) == 28 else 'other', 'input': name}
    return sf.add(sdir, name, section, len(data))


@app.route('/api/sessions/<sid>/files', methods=['GET'])
def session_files(sid):
    return jsonify({'files': sf.load(UPLOADS_DIR / sid)})


@app.route('/api/sessions/<sid>/files', methods=['POST'])
def upload_file(sid):
    """Each file goes with its own request (description 3.2, was Р-41: one big request → «HTTP 413»)."""
    from core.db import get_session
    if not get_session(DB_PATH, sid):
        return jsonify({'error': 'not_found'}), 404
    section = request.form.get('section', 'anon')
    if section not in sf.SECTIONS or 'file' not in request.files:
        return jsonify({'error': 'bad request'}), 400
    entry = _store_upload(sid, request.files['file'], section)
    if entry.get('code'):
        return jsonify(entry), 409 if entry['code'] == 'duplicate' else 400
    return jsonify(entry), 201


@app.route('/api/sessions/<sid>/files', methods=['DELETE'])
def delete_session_file(sid):
    """«Убрать из сессии»: the file and its original are deleted at once; the entries of the replacement table stay."""
    data = _body()
    name, section = data.get('input', ''), data.get('section', 'anon')
    if not _safe(name):
        return jsonify({'error': 'invalid filename'}), 400
    sdir, in_dir, out_dir = _dirs(sid, section)
    entry = sf.remove(sdir, name, section)
    if not entry:
        return jsonify({'error': 'not_found'}), 404
    (in_dir / name).unlink(missing_ok=True)
    for out in (entry.get('output'), entry.get('extra')):
        if out and _safe(out) and not any(out in (e.get('output'), e.get('extra')) for e in sf.load(sdir)):
            (out_dir / out).unlink(missing_ok=True)
            (sdir / 'restore_marks' / f'{out}.json').unlink(missing_ok=True)
    return jsonify({'ok': True})


# progress of the running file for the window: {progress id: {'stage': read / find / replace / leaks / restore}}
PROGRESS = {}
_STAGE = {'pdf_to_docx': 'read', 'process': 'read', 'known_global': 'find', 'known_session_first': 'find',
          'opf': 'find', 'rules': 'find', 'spacy': 'find', 'llm': 'find', 'propagate_orgs': 'find',
          'propagate_surnames': 'find', 'known_session': 'find', 'consolidate': 'replace', 'verify': 'leaks',
          'leak_check': 'leaks', 'restore': 'restore', 'cases': 'restore'}


@app.route('/api/progress/<pid>')
def progress(pid):
    return jsonify(PROGRESS.get(pid) or {})


def _error_code(ex) -> str:
    """Error kind for the window - what happened and what to do (description 4.6, was Р-41)."""
    import zipfile as _z
    from core.handlers import FileRejected
    if isinstance(ex, FileRejected):
        return ex.code
    if isinstance(ex, OSError) and getattr(ex, 'errno', 0) == 28:
        return 'disk'
    text = f'{type(ex).__name__} {ex}'.lower()
    if isinstance(ex, _z.BadZipFile) or any(w in text for w in ('password', 'encrypt', 'badzip', 'not a zip',
                                                                  'filedataerror', 'cannot open', 'corrupt')):
        return 'broken'
    return 'other'


def _wait_ner(use_spacy: bool, timeout=300):
    """Without the name model persons and companies would be found by the rules only (was Р-13): wait for it."""
    if not use_spacy:
        return None
    from core.anonymizer import get_ner_status
    end = time.time() + timeout
    while time.time() < end:
        s = get_ner_status()
        if s['ready']:
            return None
        if not s['loading']:
            return 'ner_failed'
        time.sleep(0.3)
    return 'ner_failed'


def _run_file(sid, name, section, use_spacy=True, lock='wait', pid=None, review=None, restore_llm=True) -> dict:
    """Processes one file of the session and records the result in the manifest."""
    from core.handlers import process_uploaded_file, IMAGE_EXT
    sdir, in_dir, out_dir = _dirs(sid, section)
    src = in_dir / name
    mode = sf.SECTIONS[section]
    files = sf.load(sdir)
    prev = sf.find(files, name, section)
    use_llm = user_settings.get()['use_llm']
    review = (use_llm if review is None else review) and section == 'anon'

    def on_stage(stage):
        if pid:
            PROGRESS[pid] = {'stage': _STAGE.get(stage, PROGRESS.get(pid, {}).get('stage', 'find')),
                             'input': name, 't': time.time()}
    if pid:
        on_stage('process')
    if os.environ.get('ANONYMIZER_SLOW'):        # window tests (UC-A2): time to press «Остановить»
        time.sleep(float(os.environ['ANONYMIZER_SLOW']))
    try:
        if prev and prev.get('output'):
            for old in (prev.get('output'), prev.get('extra')):
                if old and _safe(old):
                    (out_dir / old).unlink(missing_ok=True)
        with journal.Job(mode, on_stage=on_stage, session=sid[:8], ext=src.suffix.lower(),
                         size_kb=round(src.stat().st_size / 1024), spacy=use_spacy, review=review) as job:
            if section == 'restore' and restore_llm and use_llm:
                from core import heavy        # the case by the neural network - under the GPU queue (3.14)
                with heavy.job(lock, 'падеж при возврате данных', wait=1800):
                    r = process_uploaded_file(input_path=src, output_dir=out_dir, session_id=sid,
                                              db_path=DB_PATH, mode=mode, use_spacy=use_spacy)
            else:
                r = process_uploaded_file(input_path=src, output_dir=out_dir, session_id=sid, db_path=DB_PATH,
                                          mode=mode, use_spacy=use_spacy, restore_llm=False)
        result = {'found': r.get('entities_found', 0), 'leaks': r.get('leaks'),
                  'apply_errors': r.get('apply_errors') or 0, 'restore': r.get('restore'),
                  'converted_from': r.get('converted_from'), 'job': job.id, 'without_ner': not use_spacy}
        if section == 'anon':
            from core.db import get_occurrences
            occ = get_occurrences(DB_PATH, sid, r['output_filename'])
            result['found'] = sum(len(v) for v in occ.values()) if occ else result['found']
        if src.suffix.lower() in IMAGE_EXT:
            result['image'] = True
        if review and src.suffix.lower() not in IMAGE_EXT:
            result['llm_job'] = _start_review(sid, src, out_dir / r['output_filename'], lock)
        return sf.update(sdir, name, section, status='done', output=r['output_filename'],
                         extra=r.get('extra_output'), result=result, error=None)
    except Exception as ex:
        journal.error('process_failed', ex)
        return sf.update(sdir, name, section, status='error', error=_error_code(ex))
    finally:
        if pid:
            PROGRESS.pop(pid, None)


@app.route('/api/sessions/<sid>/process', methods=['POST'])
def process_one(sid):
    """«Обезличить (N)» / «Вернуть данные (N)»: one file per request - the window shows the progress and can stop
    between files (description 3.2, standard 2.2)."""
    data = _body()
    name, section = data.get('input', ''), data.get('section', 'anon')
    if not _safe(name) or section not in sf.SECTIONS:
        return jsonify({'error': 'bad request'}), 400
    sdir, in_dir, _ = _dirs(sid, section)
    if not (in_dir / name).exists() or not sf.find(sf.load(sdir), name, section):
        return jsonify({'error': 'not_found'}), 404
    use_spacy = data.get('use_spacy', True) is not False
    if section == 'anon':
        code = _wait_ner(use_spacy)
        if code:
            return jsonify({'code': code}), 409
    entry = _run_file(sid, name, section, use_spacy=use_spacy, lock=data.get('lock', 'wait'),
                      pid=data.get('progress'), restore_llm=data.get('llm', True) is not False)
    from core.db import get_session
    return jsonify({'file': entry, 'records': (get_session(DB_PATH, sid) or {}).get('records', 0)})


@app.route('/api/process', methods=['POST'])
def process():
    """Several files in one request (scripts, end-to-end tests): the same work as the window does file by file."""
    from core.db import get_session_mappings
    mode = request.form.get('mode', 'anonymize')
    section = 'restore' if mode == 'deanonymize' else 'anon'
    sid = request.form.get('session_id', '').strip()
    use_spacy = request.form.get('use_spacy', 'true').lower() != 'false'
    engine = request.form.get('engine') or ('accurate' if request.form.get('use_llm') == 'true' else 'fast')
    if not sid:
        return jsonify({'error': 'session_id не передан'}), 400
    files = [f for f in request.files.getlist('files') if f.filename]
    if not files:
        return jsonify({'error': 'Файлы не переданы'}), 400
    results = []
    for f in files:
        entry = _store_upload(sid, f, section)
        if entry.get('code') == 'duplicate':
            entry = sf.find(sf.load(UPLOADS_DIR / sid), entry['input'], section)
        if entry.get('code'):
            results.append({'filename': f.filename, 'status': 'error', 'error': entry['code']})
            continue
        e = _run_file(sid, entry['input'], section, use_spacy=use_spacy, lock=request.form.get('lock', 'wait'),
                      review=engine == 'accurate')
        res = e.get('result') or {}
        results.append({'filename': f.filename, 'input': e['input'], 'status': 'ok' if e['status'] == 'done' else 'error',
                        'output': e.get('output'), 'extra_output': e.get('extra'), 'error': e.get('error'),
                        'entities_found': res.get('found', 0), 'leaks': res.get('leaks'),
                        'restore': res.get('restore'), 'apply_errors': res.get('apply_errors'),
                        'converted_from': res.get('converted_from'), 'llm_job': res.get('llm_job'),
                        'job': res.get('job')})
    return jsonify({'results': results, 'session_id': sid, 'mappings': get_session_mappings(DB_PATH, sid)})


@app.route('/api/sessions/<sid>/reprocess', methods=['POST'])
def reprocess_session(sid):
    """Rebuild after every edit of the table (description 3.3, was Р-19): all finished files of «Обезличить»."""
    sdir, in_dir, _ = _dirs(sid)
    results = []
    for e in sf.load(sdir):
        if e['section'] != 'anon' or e['status'] != 'done' or not (in_dir / e['input']).exists():
            continue
        without_ner = (e.get('result') or {}).get('without_ner', False)
        n = _run_file(sid, e['input'], 'anon', use_spacy=not without_ner, review=False)
        results.append({'filename': e['input'], 'input': e['input'],
                        'status': 'ok' if n['status'] == 'done' else 'error',
                        'output': n.get('output'), 'leaks': (n.get('result') or {}).get('leaks'), 'file': n})
    return jsonify({'results': results})


# ── Replacement table ─────────────────────────────────────────────────────────
@app.route('/api/sessions/<sid>/mappings', methods=['GET'])
def session_mappings(sid):
    """Active entries; ?all=1 adds «Не маскируется» (excluded) and replaced ones. in_dict - the value is in
    «Всегда маскировать» (the record window shows the switch then)."""
    from core.db import get_session_mappings, dict_list
    rows = get_session_mappings(DB_PATH, sid, include_inactive=request.args.get('all') == '1')
    from core.cases import decline
    always = {(x['value'], x['entity_type']) for x in dict_list(DB_PATH, 'always')}
    for m in rows:
        m['in_dict'] = (m['original_form'], m['entity_type']) in always
        # «Начальная форма»: what goes into new text - a name in the nominative unless the user wrote it
        base = m['canonical_form'] or m['original_form']
        if m['entity_type'] in ('ФИО', 'FIO') and not m.get('canonical_edited'):
            try:
                base = decline(base, 'nomn')
            except Exception:
                pass
        m['base_form'] = base
    return jsonify(rows)


def _scope(data) -> str:
    return 'always' if data.get('scope') == 'always' else 'session'


@app.route('/api/sessions/<sid>/mappings/<token>', methods=['DELETE'])
def delete_mapping_route(sid, token):
    """«Не маскировать» (description 3.5): every form of the entity is excluded, the entry goes to «Не маскируется»
    (kept for earlier files); «всегда» - the value goes to the Dictionary «Никогда не маскировать»."""
    from core.db import delete_mapping, get_session_mappings, add_exclusion, dict_add
    data = _body()
    forms = [x for x in get_session_mappings(DB_PATH, sid) if x['token'] == token]
    if not forms:
        return jsonify({'error': 'not_found'}), 404
    for x in forms:
        add_exclusion(DB_PATH, sid, x['original_form'], x['entity_type'])
    delete_mapping(DB_PATH, sid, token)
    if _scope(data) == 'always':
        dict_add(DB_PATH, forms[0]['original_form'], forms[0]['entity_type'], 'never', 'unmask')
    journal.event('mapping_edit', token=token, action='exclude', scope=_scope(data))
    return jsonify({'ok': True, 'value': forms[0]['original_form'], 'token': token})


@app.route('/api/sessions/<sid>/mappings/<token>/mask', methods=['POST'])
def mask_again_route(sid, token):
    """«Вернуть маску» in «Не маскируется» and «Вернуть» after «Не маскировать» (was Р-22): the same token again."""
    from core.db import unexclude_token, dict_remove
    forms = unexclude_token(DB_PATH, sid, token)
    if not forms:
        return jsonify({'error': 'not_found'}), 404
    left_dict = [v for v, t in forms if dict_remove(DB_PATH, v, t, 'never')]
    journal.event('mapping_edit', token=token, action='mask_again')
    return jsonify({'ok': True, 'token': token, 'removed_from_dictionary': len(left_dict)})


@app.route('/api/sessions/<sid>/mappings/<token>', methods=['PATCH'])
def update_mapping_route(sid, token):
    from core.db import (get_session_mappings, add_exclusion, get_or_create_token, dict_add, update_mapping,
                         retire_mapping, remove_exclusions_for)
    data = _body()
    old = next((x for x in get_session_mappings(DB_PATH, sid) if x['token'] == token), None)
    if not old:
        return jsonify({'error': 'Token not found'}), 404
    new_original = (data.get('original_form', old['original_form'])).strip()
    new_canonical = (data.get('canonical_form', old['canonical_form'])).strip() or new_original
    new_type = (data.get('entity_type', old['entity_type'])).strip()
    choice = data.get('on_duplicate')       # None → ask; 'merge' / 'separate'
    if (new_original != old['original_form'] or new_type != old['entity_type']) and not choice:
        dup = _duplicate_of(sid, new_original, new_type, exclude=token)
        if dup:
            return jsonify({'duplicate': dup}), 409
    if _scope(data) == 'always':
        dict_add(DB_PATH, new_original, new_type, 'always', 'manual')
    if new_original == old['original_form'] and new_type == old['entity_type']:
        # only the initial form changed: same token, nothing to re-detect
        if new_canonical != old['canonical_form']:
            update_mapping(DB_PATH, sid, token, {'canonical_form': new_canonical})
        journal.event('mapping_edit', token=token, action='canonical')
        return jsonify({'ok': True, 'new_token': token})
    # value or type changed: the old token is retired (kept to restore earlier files)
    if new_original != old['original_form']:
        add_exclusion(DB_PATH, sid, old['original_form'], old['entity_type'])
    retire_mapping(DB_PATH, sid, token)
    remove_exclusions_for(DB_PATH, sid, new_original, new_type)
    edited = bool(old.get('canonical_edited')) or new_canonical != old['canonical_form']
    new_token = get_or_create_token(DB_PATH, sid, new_original, new_canonical, new_type,
                                    canonical_edited=edited, manual=True, separate=choice == 'separate')
    journal.event('mapping_edit', token=token, action='replace', new_token=new_token)
    return jsonify({'ok': True, 'new_token': new_token})


@app.route('/api/sessions/<sid>/mappings', methods=['POST'])
def add_mapping_route(sid):
    """«+ Добавить вручную» - also before the first file (was Р-20) - and «Замаскировать» in the preview.
    In the Dictionary - only with «всегда» (was Р-24: everything went there silently)."""
    from core.db import get_or_create_token, dict_add, remove_exclusions_for
    data = _body()
    original = (data.get('original_form') or '').strip()
    canonical = (data.get('canonical_form') or '').strip()
    entity_type = (data.get('entity_type') or '').strip()
    if not entity_type:
        return jsonify({'error': 'entity_type обязателен'}), 400
    edited = bool(canonical) and bool(original) and canonical != original
    canonical = canonical or original
    original = original or canonical
    if not original:
        return jsonify({'error': 'empty', 'code': 'empty'}), 400
    choice = data.get('on_duplicate')
    if not choice:
        dup = _duplicate_of(sid, original, entity_type)
        if dup:
            return jsonify({'duplicate': dup}), 409
    remove_exclusions_for(DB_PATH, sid, original, entity_type)    # added by hand: mask it again
    token = get_or_create_token(DB_PATH, sid, original, canonical, entity_type,
                                canonical_edited=edited, manual=True, separate=choice == 'separate')
    if _scope(data) == 'always':
        dict_add(DB_PATH, original, entity_type, 'always', 'manual')
    journal.event('mapping_edit', token=token, action='add', scope=_scope(data))
    return jsonify({'token': token, 'original_form': original, 'canonical_form': canonical,
                    'entity_type': entity_type}), 201


def _duplicate_of(sid, original, entity_type, exclude=None):
    """«Такое значение уже есть: [ADR_2] …»: the active entry this value already belongs to, or None."""
    from core.db import find_same_entity, get_session_mappings
    tok = find_same_entity(DB_PATH, sid, original, entity_type)
    if not tok or tok == exclude:
        return None
    forms = [m for m in get_session_mappings(DB_PATH, sid) if m['token'] == tok]
    if not forms:
        return None
    return {'token': tok, 'value': forms[0]['canonical_form'] or forms[0]['original_form'],
            'entity_type': forms[0]['entity_type']}


@app.route('/api/sessions/<sid>/mappings/<token>/alias', methods=['POST'])
def add_alias_route(sid, token):
    """«Другое написание…»: another written form of this entity gets the same token."""
    from core.db import get_session_mappings, add_alias, dict_add
    data = _body()
    value = (data.get('original_form') or '').strip()
    m = next((x for x in get_session_mappings(DB_PATH, sid) if x['token'] == token), None)
    if not m or not value:
        return jsonify({'error': 'empty', 'code': 'empty'}), 400
    add_alias(DB_PATH, sid, token, value, m['entity_type'])
    if _scope(data) == 'always':
        dict_add(DB_PATH, value, m['entity_type'], 'always', 'manual')
    journal.event('mapping_edit', token=token, action='alias')
    return jsonify({'ok': True, 'token': token})


@app.route('/api/sessions/<sid>/mappings/<token>/merge', methods=['POST'])
def merge_mapping(sid, token):
    """«Объединить с…»: every form of `token` moves under `into`; the old token is kept for earlier files."""
    from core.db import get_session_mappings, merge_tokens
    into = _body().get('into')
    maps = get_session_mappings(DB_PATH, sid)
    if not any(m['token'] == token for m in maps) or not any(m['token'] == into for m in maps) or into == token:
        return jsonify({'error': 'not_found'}), 404
    merge_tokens(DB_PATH, sid, token, into)
    journal.event('mapping_edit', token=token, action='merge', into=into)
    return jsonify({'ok': True, 'mappings': get_session_mappings(DB_PATH, sid)})


@app.route('/api/sessions/<sid>/duplicates', methods=['GET'])
def session_duplicates(sid):
    """«Проверка дублей»: pairs proposed to a human (never numbers), minus «Это разные»."""
    from core.db import get_session_mappings, get_distinct_pairs
    from core.entities import review_pairs
    return jsonify(review_pairs(get_session_mappings(DB_PATH, sid), get_distinct_pairs(DB_PATH, sid)))


@app.route('/api/sessions/<sid>/distinct', methods=['POST'])
def mark_distinct_route(sid):
    """«Это разные»: the pair is remembered and not proposed or merged again."""
    from core.db import mark_distinct
    data = _body()
    for a, b in data.get('pairs') or [[data.get('a'), data.get('b')]]:
        if a and b and a != b:
            mark_distinct(DB_PATH, sid, a, b)
            journal.event('mapping_edit', token=a, action='distinct', other=b)
    return jsonify({'ok': True})


# ── Dictionary (description 3.11) ─────────────────────────────────────────────
@app.route('/api/dictionary', methods=['GET'])
def dictionary():
    from core.db import dict_list
    st = user_settings.get()
    return jsonify({'always': dict_list(DB_PATH, 'always'), 'never': dict_list(DB_PATH, 'never'),
                    'migrated': st['dict_migrated'] if not st['dict_note_hidden'] else 0})


@app.route('/api/dictionary', methods=['POST'])
def dictionary_add():
    """«+ Добавить» and «Вернуть» after ✕: into one list, out of the other."""
    from core.db import dict_add
    d = _body()
    which = d.get('list') if d.get('list') in ('always', 'never') else 'always'
    value, etype = (d.get('value') or '').strip(), (d.get('entity_type') or '').strip()
    if not value or not etype:
        return jsonify({'error': 'empty', 'code': 'empty'}), 400
    moved = dict_add(DB_PATH, value, etype, which, d.get('source') or 'manual', d.get('added_at'))
    journal.event('dictionary', action='add', list=which, moved=moved)
    return jsonify({'ok': True, 'moved': moved}), 201


@app.route('/api/dictionary/remove', methods=['POST'])
def dictionary_remove():
    from core.db import dict_remove
    d = _body()
    which = d.get('list') if d.get('list') in ('always', 'never') else 'always'
    row = dict_remove(DB_PATH, d.get('value') or '', d.get('entity_type') or '', which)
    journal.event('dictionary', action='remove', list=which)
    return jsonify({'ok': bool(row), 'row': row})


@app.route('/api/dictionary/clear', methods=['POST'])
def dictionary_clear():
    from core.db import dict_clear
    which = _body().get('list')
    if which not in ('always', 'never'):
        return jsonify({'error': 'bad request'}), 400
    n = dict_clear(DB_PATH, which)
    journal.event('dictionary', action='clear', list=which, count=n)
    return jsonify({'ok': True, 'removed': n})


@app.route('/api/dictionary/note-hide', methods=['POST'])
def dictionary_note_hide():
    user_settings.update({'dict_note_hidden': True})
    return jsonify({'ok': True})


# ── Background check by the neural network ────────────────────────────────────
REVIEWS = {}     # job id → state; in memory, a review lives as long as the app runs
CANCEL = {}      # job id → threading.Event: «Отменить ожидание» / «Остановить»


def _start_review(sid, src: Path, out: Path, lock: str = 'wait') -> str:
    """The file is already downloadable; the local neural network checks the found values and the suspicious lines
    in a thread and proposes additions / removals (description 3.15)."""
    jid = uuid.uuid4().hex[:12]
    REVIEWS[jid] = {'id': jid, 'status': 'running', 'session': sid, 'input': src.name,
                    'output': out.name, 'started': time.time(), 'updated': time.time()}
    CANCEL[jid] = threading.Event()

    def work():
        from core.extract import extract_text
        from core.db import get_session_mappings
        from core.handlers import CONVERT_EXT, _convert_to_docx
        from core import heavy
        from core.llm_review import review
        st = REVIEWS[jid]

        def waiting(name):          # in the GPU queue: «Нейросеть ждёт видеокарту: «Herald»»
            st.update(status='wait_lock', holder=name, updated=time.time())
        try:
            with heavy.job(lock, 'проверка нейросетью', wait=1800, on_wait=waiting, cancel=CANCEL[jid]), \
                    journal.Job('llm_review', session=sid[:8], ext=src.suffix.lower()):
                if st.get('cancel'):
                    raise heavy.Busy('cancelled')
                if st['status'] == 'wait_lock':
                    st.update(status='running', started=time.time())
                orig_path = src
                if src.suffix.lower() in CONVERT_EXT:
                    import tempfile
                    orig_path = _convert_to_docx(src, Path(tempfile.mkdtemp()))
                original = extract_text(orig_path, meta=False)
                masked = extract_text(out, meta=False)
                st['proposal'] = review(original, masked, get_session_mappings(DB_PATH, sid))
                st['status'] = 'done'
        except Exception as ex:
            journal.error('llm_review_failed', ex)
            if st.get('cancel'):
                st['status'] = 'cancelled'
            else:
                st['status'] = 'error'
                st['error'] = 'busy' if isinstance(ex, heavy.Busy) else 'unavailable'
        st['updated'] = time.time()

    threading.Thread(target=work, daemon=True).start()
    return jid


@app.route('/api/llm-jobs/<jid>')
def llm_job_status(jid):
    from core.db import get_session_mappings
    st = REVIEWS.get(jid)
    if not st:
        return jsonify({'error': 'lost', 'status': 'lost'}), 404
    out = {k: v for k, v in st.items() if k != 'proposal'}
    if st['status'] == 'running' and time.time() - st['started'] > 300:
        out['status'] = 'stale'
    out['elapsed'] = round(time.time() - st['started'], 1)
    p = st.get('proposal')
    if p:
        vals = {m['token']: m for m in get_session_mappings(DB_PATH, st['session'])}
        out['proposal'] = {
            'add': p['add'],
            'remove': [dict(x, value=vals.get(x['token'], {}).get('original_form', '')) for x in p['remove']
                       if x['token'] in vals],
            'merge': [dict(x, value=vals.get(x['token'], {}).get('original_form', ''),
                           into_value=vals.get(x['into'], {}).get('original_form', ''))
                      for x in p['merge'] if x['token'] in vals and x['into'] in vals],
            'seconds': p.get('seconds'), 'checked': p.get('checked')}
    return jsonify(out)


@app.route('/api/llm-jobs/<jid>/apply', methods=['POST'])
def llm_job_apply(jid):
    """«Применить»: the ticked proposals, then the file is rebuilt."""
    from core.db import get_session_mappings, add_exclusion, delete_mapping, get_or_create_token, merge_tokens
    st = REVIEWS.get(jid)
    if not st or st.get('status') != 'done':
        return jsonify({'error': 'no proposal'}), 400
    data = _body()
    sid, p = st['session'], st['proposal']
    remove = set(data.get('remove', []))
    for m in get_session_mappings(DB_PATH, sid):
        if m['token'] in remove:
            add_exclusion(DB_PATH, sid, m['original_form'], m['entity_type'])
    for t in remove:
        delete_mapping(DB_PATH, sid, t)
    for i in data.get('add', []):
        if 0 <= i < len(p['add']):
            a = p['add'][i]
            get_or_create_token(DB_PATH, sid, a['value'], a['value'], a['type'], manual=True)
    for x in p['merge']:
        if x['token'] in set(data.get('merge', [])):
            merge_tokens(DB_PATH, sid, x['token'], x['into'])
    journal.event('llm_review_applied', removed=len(remove), added=len(data.get('add', [])),
                  merged=len(data.get('merge', [])))
    entry = _run_file(sid, st['input'], 'anon', review=False)
    entry = sf.update(UPLOADS_DIR / sid, st['input'], 'anon',
                      result=dict(entry.get('result') or {}, llm_applied=True))
    st['status'] = 'applied'
    return jsonify({'ok': True, 'output': entry.get('output'), 'file': entry,
                    'mappings': get_session_mappings(DB_PATH, sid)})


# ── Preview ───────────────────────────────────────────────────────────────────
@app.route('/api/sessions/<sid>/preview/<path:filename>')
def preview_file(sid, filename):
    """Text with tokens, forms, places, the leak check, duplicate pairs (description 3.6); for a file with returned
    data - what was inserted and how (3.8)."""
    from core.extract import extract_text
    from core.handlers import IMAGE_EXT, leak_check, load_restore_marks
    from core.db import get_session_mappings, get_occurrences, get_distinct_pairs
    from core.entities import review_pairs
    if not _safe(filename):
        return jsonify({'error': 'invalid filename'}), 400
    path = UPLOADS_DIR / sid / 'output' / filename
    if not path.exists():
        return jsonify({'error': 'not_found'}), 404
    maps = get_session_mappings(DB_PATH, sid)
    forms = {}
    for m in maps:
        forms.setdefault(m['token'], {'type': m['entity_type'], 'forms': []})['forms'].append(m['original_form'])
    ext = path.suffix.lower()
    entry = sf.by_output(sf.load(UPLOADS_DIR / sid), filename) or {}
    out = {'filename': filename, 'forms': forms, 'kind': 'text', 'pages': 0, 'text': '',
           'input': entry.get('input'), 'section': entry.get('section'),
           'occurrences': get_occurrences(DB_PATH, sid, filename)}
    marks = load_restore_marks(path.parent, filename)
    if marks is not None:
        out['restore'] = {'marks': marks,
                          'exact': sum(1 for m, _ in marks if m == 'exact'),
                          'case': sum(1 for m, _ in marks if m == 'case'),
                          'unknown': sum(1 for m, _ in marks if m == 'unknown')}
    if ext in ('.docx', '.docm'):
        out['kind'] = 'docx'
    if ext in IMAGE_EXT:
        out['kind'] = 'image'
    elif ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True)
        out['kind'] = 'xlsx'
        sheets = []
        for ws in wb.worksheets:
            rows = [['' if v is None else str(v) for v in row[:40]]
                    for row in ws.iter_rows(max_row=801, values_only=True)]
            cut = len(rows) > 800 or (ws.max_column or 0) > 40
            sheets.append({'name': ws.title, 'rows': rows[:800], 'truncated': cut})
        out['sheets'] = sheets
        out['text'] = extract_text(path, meta=False)[:300000]
        out['leaks'] = None if marks is not None else leak_check(path, sid, DB_PATH)
    else:
        try:
            out['text'] = extract_text(path, meta=False)[:300000]
        except Exception as ex:
            journal.error('preview_failed', ex)
            out['text'], out['error'] = '', 'unreadable'
        if ext == '.pdf':
            import pymupdf
            out['kind'] = 'pdf'
            out['pages'] = len(pymupdf.open(str(path)))
        out['leaks'] = None if marks is not None else leak_check(path, sid, DB_PATH)
    out['similar'] = review_pairs(maps, get_distinct_pairs(DB_PATH, sid))
    return jsonify(out)


@app.route('/api/sessions/<sid>/page/<path:filename>/<int:page>.png')
def preview_pdf_page(sid, filename, page):
    import pymupdf
    if not _safe(filename):
        return jsonify({'error': 'invalid filename'}), 400
    doc = pymupdf.open(str(UPLOADS_DIR / sid / 'output' / filename))
    if not 0 <= page < len(doc):
        return jsonify({'error': 'page'}), 404
    return send_file(io.BytesIO(doc[page].get_pixmap(dpi=110).tobytes('png')), mimetype='image/png')


@app.route('/api/sessions/<sid>/restore-text', methods=['POST'])
def restore_text_route(sid):
    """«Вернуть данные → Текст» (question 1): the pasted reply is not written anywhere."""
    from core.anonymizer import restore_pasted
    from core.db import get_session
    text = _body().get('text') or ''
    if not get_session(DB_PATH, sid):
        return jsonify({'error': 'not_found'}), 404
    if not text.strip():
        return jsonify({'error': 'empty', 'code': 'empty'}), 400
    data = _body()
    llm = user_settings.get()['use_llm'] and data.get('llm', True) is not False
    with journal.Job('restore_text', session=sid[:8], chars=len(text), llm=llm):
        if llm:
            from core import heavy
            with heavy.job(data.get('lock', 'wait'), 'падеж при возврате данных', wait=1800):
                r = restore_pasted(text[:500000], DB_PATH, sid, use_llm=True)
        else:
            r = restore_pasted(text[:500000], DB_PATH, sid, use_llm=False)
    return jsonify(r)


# ── Downloads: only names with tokens (правка 6) ──────────────────────────────
@app.route('/api/download/<sid>/<filename>')
def download_file(sid, filename):
    if not _safe(filename):
        return jsonify({'error': 'invalid filename'}), 400
    file_path = UPLOADS_DIR / sid / 'output' / filename
    if not file_path.exists():
        return jsonify({'error': 'not_found'}), 404
    journal.event('download', session=sid[:8], ext=file_path.suffix.lower(),
                  size_kb=round(file_path.stat().st_size / 1024))
    return send_file(str(file_path), as_attachment=True, download_name=filename)


@app.route('/api/download-all/<sid>')
def download_all(sid):
    """«Скачать все (ZIP)» - only the files of one section (was Р-40); the archive name has no user text."""
    from core.db import get_session
    section = request.args.get('section', 'anon')
    sdir, _, out_dir = _dirs(sid)
    names = []
    for e in sf.load(sdir):
        if e['section'] == section and e['status'] == 'done':
            names += [n for n in (e.get('output'), e.get('extra')) if n and (out_dir / n).exists()]
    if not names:
        return jsonify({'error': 'not_found'}), 404
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for n in dict.fromkeys(names):
            zf.write(str(out_dir / n), n)
    buf.seek(0)
    s = get_session(DB_PATH, sid) or {}
    day = (s.get('created_at') or '')[8:10] + '.' + (s.get('created_at') or '')[5:7]
    title = 'Обезличенные файлы' if section == 'anon' else 'Данные возвращены'
    journal.event('download_all', session=sid[:8], section=section, files=len(names))
    return send_file(buf, mimetype='application/zip', as_attachment=True, download_name=f'{title} - {day}.zip')


# ── Journal, data folder ──────────────────────────────────────────────────────
def _open_folder(path: Path):
    import subprocess
    opener = 'open' if sys.platform == 'darwin' else ('explorer' if sys.platform == 'win32' else 'xdg-open')
    subprocess.Popen([opener, str(path)])


@app.route('/api/logs')
def logs_info():
    return jsonify({'dir': str(journal.LOG_DIR), 'verbose': journal.VERBOSE, 'stale_jobs': journal.stale_jobs()})


@app.route('/api/jobs/recent')
def recent_jobs():
    """«Последние задачи»: the last 20 jobs - time, file type, mode, duration, status; no names, no values."""
    return jsonify(journal.recent_jobs(20))


@app.route('/api/logs/open', methods=['POST'])
def logs_open():
    if os.environ.get('ANONYMIZER_NO_OPEN'):
        return jsonify({'ok': True, 'dir': str(journal.LOG_DIR)})
    try:
        _open_folder(journal.LOG_DIR)
    except OSError:
        return jsonify({'error': 'open_failed', 'dir': str(journal.LOG_DIR)}), 500
    return jsonify({'ok': True, 'dir': str(journal.LOG_DIR)})


@app.route('/api/data/open', methods=['POST'])
def data_open():
    if os.environ.get('ANONYMIZER_NO_OPEN'):
        return jsonify({'ok': True, 'dir': str(DATA_DIR)})
    try:
        _open_folder(DATA_DIR)
    except OSError:
        return jsonify({'error': 'open_failed', 'dir': str(DATA_DIR)}), 500
    return jsonify({'ok': True, 'dir': str(DATA_DIR)})


@app.route('/api/logs/verbose', methods=['POST'])
def logs_verbose():
    """«Режим диагностики»: values in the journal until the restart (description 3.12)."""
    journal.set_verbose(bool(_body().get('on')))
    return jsonify({'verbose': journal.VERBOSE})


@app.route('/api/shutdown', methods=['POST'])
def shutdown():
    def _stop():
        time.sleep(0.5)
        os._exit(0)
    threading.Thread(target=_stop, daemon=True).start()
    return jsonify({'ok': True})


# ── Launch ────────────────────────────────────────────────────────────────────
def _already_running(port) -> bool:
    import json
    import urllib.request
    try:
        with urllib.request.urlopen(f'http://127.0.0.1:{port}/api/info', timeout=1) as r:
            return json.loads(r.read()).get('name') == 'Anonymizer'
    except Exception:
        return False


def main():
    import socket
    with socket.socket() as sock:
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)     # the previous run's TIME_WAIT is not «busy»
        try:
            sock.bind(('127.0.0.1', PORT))
        except OSError:
            if _already_running(PORT):
                print(f'Anonymizer уже работает: http://127.0.0.1:{PORT}', flush=True)
                return
            print(f'Порт {PORT} занят другой программой - Anonymizer не запущен', flush=True)
            sys.exit(1)
    print('Anonymizer ' + APP_VERSION, flush=True)
    print(f'http://127.0.0.1:{PORT}', flush=True)
    print('Data:', DATA_DIR, flush=True)
    app.run(host='127.0.0.1', port=PORT, debug=False, use_reloader=False, threaded=True)


if __name__ == '__main__':
    main()
